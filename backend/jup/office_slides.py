"""PowerPoint decks for the assistant (Track O, O3 1/3).

Like `jup.office_doc` and `jup.office_sheet`, the assistant never writes a file itself. It sends a
**spec** (a title, a cover and a list of slides, each naming one of the template's layouts and
filling that layout's fields) and this module builds the .pptx on the server from the JUP template
(Template presentasi JUP v1.5-r2, in `office_templates/pptx/`, described by its `manifest.json`).

**Rules the model cannot override** (each tested in `tests/test_jup_office_slides.py`):

- **Every slide stays on the master.** A slide is a template layout with its placeholders filled;
  the renderer never draws a free-floating shape. Tables and charts go into the layout's table and
  chart placeholders and take their styling from the template's own example slides, which are
  then removed, so no example text ("Data ilustratif") reaches the output.
- **Nothing is shrunk to fit.** Each text box in the template has a fixed size and font size; the
  renderer measures the text against it and refuses with a message the model can act on.
- **Charts stay editable:** native PowerPoint charts with their data in the embedded workbook.
- The deck always opens with a cover the renderer builds, and the status is always **Draf**. The
  classification follows `office_doc.classification_for` (INTERNAL by default, never PUBLIK, never
  below the caller's floor) and is written in every footer the template has.
- Limits are refused with a message, never silently truncated.

The template writes `type="subtitle"` on most subtitle placeholders, where OOXML (and python-pptx)
spell it `subTitle`; the file is loaded and that one attribute corrected in memory, so the
checked-in template stays byte-identical to the one JUP issued.

Pure: no app imports, no network, no disk writes. `jup.office_tool` stores the bytes.
"""

from __future__ import annotations

import copy
import datetime as dt
import io
import json
import math
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from jup.office_doc import MAX_TITLE, STATUS, OfficeDocError, _text, classification_for
from jup.office_sheet import JAKARTA, _date, _number

TEMPLATE_DIR = Path(__file__).resolve().parent / 'office_templates' / 'pptx'
PPTX_MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

MAX_SLIDES = 30  # including the cover
MAX_NOTES = 3_000
MAX_FIELD = 1_000  # any one text before measuring; the layout's box is the real limit
MAX_TABLE_COLS = 6
MAX_CATEGORIES = 12
MAX_SERIES = 3
MAX_BYTES = 10 * 1024 * 1024

TABLE_TYPES = ('text', 'number', 'idr', 'percent', 'date')
_NUMERIC = frozenset({'number', 'idr', 'percent'})
COVERS = ('institusional', 'korporat')
WEBSITE_IDX = 96
FOOTER_IDX = 98

# Average advance of a character as a share of the font size, measured on the template's own
# renders with a margin: Space Grotesk prose runs about 0.48-0.54 em, Space Mono is 0.6 em.
_EM = {'Space Grotesk': 0.56, 'Space Mono': 0.62}
_LINE = 1.2  # line height as a share of the font size
_EMU_PER_PT = 12_700
_BOLD = re.compile(r'\*\*(.+?)\*\*')
_PALETTE = ('00A550', '004999', '536679')  # JUP green, Polynesian blue, slate (the template's chart colours)
_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


@dataclass
class SlidesRendered:
    data: bytes
    filename: str
    title: str
    classification: str
    status: str = STATUS
    slides: int = 0
    charts: int = 0
    tables: int = 0
    unfilled: list[str] = field(default_factory=list)
    content_type: str = PPTX_MIME


@dataclass(frozen=True)
class Box:
    """What one template text box holds: lines of about `chars` characters each."""

    chars: int
    lines: int

    def describe(self) -> str:
        return f'{self.lines} line{"s" if self.lines > 1 else ""} of about {self.chars} characters'


# -- the template ------------------------------------------------------------------------------


@lru_cache(maxsize=1)
def manifest() -> dict:
    return json.loads((TEMPLATE_DIR / 'manifest.json').read_text(encoding='utf-8'))


def template_path() -> Path:
    return TEMPLATE_DIR / manifest()['file']


@lru_cache(maxsize=1)
def _template_bytes() -> bytes:
    return template_path().read_bytes()


def _open_template():
    """The template as a fresh Presentation, with the `subtitle` placeholder type corrected."""
    prs = Presentation(io.BytesIO(_template_bytes()))
    for part in prs.part.package.iter_parts():
        element = getattr(part, '_element', None)
        if element is None:
            continue
        for ph in element.iter(qn('p:ph')):
            if ph.get('type') == 'subtitle':
                ph.set('type', 'subTitle')
    return prs


def _layouts_by_name(prs) -> dict:
    return {layout.name: layout for master in prs.slide_masters for layout in master.slide_layouts}


def _ph_by_idx(container) -> dict[int, Any]:
    out = {}
    for shape in container.shapes._spTree.iterchildren():
        ph = shape.find(f'.//{{{_P}}}nvPr/{{{_P}}}ph')
        if ph is not None:
            out[int(ph.get('idx', '0'))] = shape
    return out


def _box(ph_element) -> Box:
    """The capacity of a layout text box, from its size and the font size of its prompt."""
    ext = ph_element.find(f'.//{{{_A}}}xfrm/{{{_A}}}ext')
    width_pt = int(ext.get('cx')) / _EMU_PER_PT
    height_pt = int(ext.get('cy')) / _EMU_PER_PT
    sizes = [int(e.get('sz')) / 100 for e in ph_element.iter(qn('a:defRPr'), qn('a:rPr')) if e.get('sz')]
    size = max(sizes) if sizes else 18.0
    faces = [e.get('typeface') for e in ph_element.iter(qn('a:latin'))]
    em = _EM.get(faces[0] if faces else '', 0.56)
    return Box(chars=max(1, math.floor(width_pt / (em * size))), lines=max(1, math.floor(height_pt / (_LINE * size))))


@lru_cache(maxsize=1)
def capacities() -> dict[str, dict[int, Box]]:
    """Every working layout's text boxes by placeholder idx, read once from the template."""
    layouts = _layouts_by_name(_open_template())
    names = [c['layout'] for c in manifest()['covers'].values()] + [s['layout'] for s in manifest()['layouts'].values()]
    out = {}
    for name in names:
        out[name] = {
            idx: _box(el) for idx, el in _ph_by_idx(layouts[name]).items() if el.find(f'.//{{{_P}}}txBody') is not None
        }
    return out


@lru_cache(maxsize=1)
def section_labels() -> dict[str, str]:
    """Each content layout's own label above the title (PEMBAHASAN, RINGKASAN, …)."""
    layouts = _layouts_by_name(_open_template())
    return {spec['layout']: _layout_text(layouts[spec['layout']], 2) for spec in manifest()['layouts'].values()}


def _layout_text(layout, idx: int) -> str:
    el = _ph_by_idx(layout).get(idx)
    if el is None:
        return ''
    return '\n'.join(''.join(t.text or '' for t in p.iter(qn('a:t'))) for p in el.iter(qn('a:p'))).strip()


# -- measuring ---------------------------------------------------------------------------------


def _plain(text: str) -> str:
    return _BOLD.sub(r'\1', text)


def lines_needed(text: str, chars: int) -> int:
    """Lines `text` takes when wrapped at `chars` characters, word by word, as PowerPoint does."""
    total = 0
    for paragraph in _plain(text).split('\n'):
        words = paragraph.split()
        if not words:
            total += 1
            continue
        lines, width = 1, 0
        for word in words:
            length = len(word)
            if width == 0:
                lines += (length - 1) // chars
                width = length % chars or chars
            elif width + 1 + length <= chars:
                width += 1 + length
            else:
                lines += 1 + (length - 1) // chars
                width = length % chars or chars
        total += lines
    return total


def _fit(text: str, box: Box, where: str) -> None:
    needed = lines_needed(text, box.chars)
    if needed > box.lines:
        raise OfficeDocError(
            f'{where} does not fit: it needs about {needed} lines and the template box holds {box.describe()}. '
            'Shorten it (the font is never made smaller), or move the detail to the speaker notes.'
        )


# -- spec validation ---------------------------------------------------------------------------


def _field_text(value: Any, where: str, as_list: bool, spaced: bool = False) -> str:
    """A field's text; a list becomes one paragraph per item, with a blank line between items where
    the template's own prompt has one."""
    if as_list and isinstance(value, list):
        if not value:
            raise OfficeDocError(f'{where} is an empty list.')
        items = [_text(item, f'An item in {where}', MAX_FIELD) for item in value]
    else:
        items = [_text(value, where, MAX_FIELD)]
        if not (as_list and spaced):
            return items[0]
        items = [line.strip() for line in items[0].split('\n') if line.strip()]
    return ('\n\n' if spaced else '\n').join(items)


def _format_number(value: float | int, decimals: int) -> str:
    """Indonesian style: 1.250.000 and 12,5."""
    text = f'{value:,.{decimals}f}'
    return text.replace(',', '\x00').replace('.', ',').replace('\x00', '.')


def _decimals(values: list[Any]) -> int:
    floats = [v for v in values if isinstance(v, float) and not v.is_integer()]
    return 0 if not floats else min(2, max(len(repr(v).split('.')[1]) for v in floats))


def _table(spec: Any, where: str, max_rows: int) -> dict:
    if not isinstance(spec, dict):
        raise OfficeDocError(f'{where} needs "table": {{"columns": [{{"header", "type"}}], "rows": [[...]]}}.')
    unknown = sorted(set(spec) - {'columns', 'rows'})
    if unknown:
        raise OfficeDocError(f'{where} table has unknown keys: {", ".join(unknown)}.')
    columns = spec.get('columns')
    if not isinstance(columns, list) or not 2 <= len(columns) <= MAX_TABLE_COLS:
        raise OfficeDocError(f'{where} table needs 2 to {MAX_TABLE_COLS} columns.')
    cols = []
    for c, column in enumerate(columns, start=1):
        if not isinstance(column, dict):
            raise OfficeDocError(f'{where} table column {c} must be an object with "header" and "type".')
        kind = str(column.get('type') or 'text').strip().lower()
        if kind not in TABLE_TYPES:
            raise OfficeDocError(
                f'{where} table column {c} has type "{column.get("type")}"; use one of: {", ".join(TABLE_TYPES)}.'
            )
        cols.append({'header': _text(column.get('header'), f'{where} table column {c} header', 60), 'type': kind})
    rows = spec.get('rows')
    if not isinstance(rows, list) or not rows:
        raise OfficeDocError(f'{where} table needs at least one row.')
    if len(rows) > max_rows:
        raise OfficeDocError(
            f'{where} table has {len(rows)} rows; this layout holds {max_rows} under the header. Split it over '
            'two slides, or attach the full table as a spreadsheet with create_spreadsheet.'
        )
    raw = []
    for r, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) > len(cols):
            raise OfficeDocError(f'{where} table row {r} must be a list of at most {len(cols)} cells.')
        padded = row + [None] * (len(cols) - len(row))
        cells = []
        for value, column in zip(padded, cols):
            here = f'{where} table row {r}, column "{column["header"]}"'
            if value is None or (isinstance(value, str) and not value.strip()):
                cells.append(None)
            elif column['type'] in _NUMERIC:
                cells.append(_number(value, here))
            elif column['type'] == 'date':
                cells.append(_date(value, here))
            else:
                cells.append(_text(value if isinstance(value, str) else str(value), here, MAX_FIELD))
        raw.append(cells)
    text_rows = []
    for c, column in enumerate(cols):
        decimals = _decimals([row[c] for row in raw])
        for r, row in enumerate(raw):
            value = row[c]
            if r == len(text_rows):
                text_rows.append([])
            if value is None:
                text_rows[r].append('')
            elif column['type'] == 'number':
                text_rows[r].append(_format_number(value, decimals))
            elif column['type'] == 'idr':
                text_rows[r].append('Rp ' + _format_number(value, decimals))
            elif column['type'] == 'percent':
                text_rows[r].append(
                    _format_number(value, max(decimals, 1) if not float(value).is_integer() else 0) + '%'
                )
            elif column['type'] == 'date':
                text_rows[r].append(value.strftime('%d/%m/%Y'))
            else:
                text_rows[r].append(value)
    return {'columns': cols, 'rows': text_rows}


def _chart(spec: Any, where: str) -> dict:
    if not isinstance(spec, dict):
        raise OfficeDocError(f'{where} needs "chart": {{"categories": [...], "series": [{{"name", "values"}}]}}.')
    unknown = sorted(set(spec) - {'categories', 'series'})
    if unknown:
        raise OfficeDocError(f'{where} chart has unknown keys: {", ".join(unknown)}.')
    categories = spec.get('categories')
    if not isinstance(categories, list) or not 2 <= len(categories) <= MAX_CATEGORIES:
        raise OfficeDocError(f'{where} chart needs 2 to {MAX_CATEGORIES} categories.')
    labels = [_text(c if isinstance(c, str) else str(c), f'A category in {where}', 24) for c in categories]
    series = spec.get('series')
    if not isinstance(series, list) or not 1 <= len(series) <= MAX_SERIES:
        raise OfficeDocError(f'{where} chart needs 1 to {MAX_SERIES} series.')
    out = []
    for s, item in enumerate(series, start=1):
        if not isinstance(item, dict):
            raise OfficeDocError(f'{where} chart series {s} must be {{"name", "values"}}.')
        name = _text(item.get('name'), f'{where} chart series {s} name', 40)
        values = item.get('values')
        if not isinstance(values, list) or len(values) != len(labels):
            raise OfficeDocError(f'{where} chart series "{name}" needs exactly {len(labels)} values, one per category.')
        out.append(
            {
                'name': name,
                'values': [
                    None if v is None else _number(v, f'{where} chart series "{name}" value {i}')
                    for i, v in enumerate(values, 1)
                ],
            }
        )
    return {'categories': labels, 'series': out}


def _cover(spec: Any, title: str) -> dict:
    spec = spec or {}
    if not isinstance(spec, dict):
        raise OfficeDocError('"cover" must be an object such as {"style": "institusional", "kind": "RAPAT DIREKSI"}.')
    unknown = sorted(set(spec) - {'style', 'kind', 'unit', 'date'})
    if unknown:
        raise OfficeDocError(f'The cover has unknown keys: {", ".join(unknown)}. Use style, kind, unit, date.')
    style = str(spec.get('style') or 'institusional').strip().lower()
    if style not in COVERS:
        raise OfficeDocError(
            f'Unknown cover style "{spec.get("style")}". Use "institusional" (light) or "korporat" (dark).'
        )
    return {
        'style': style,
        'title': title,
        'kind': _text(spec.get('kind'), 'The cover "kind"', 60, required=False).upper(),
        'unit': _text(spec.get('unit'), 'The cover "unit"', 80, required=False),
        'date': _text(spec.get('date'), 'The cover "date"', 40, required=False),
    }


def validate_slides(slides: Any) -> list[dict]:
    """The slides, checked against the manifest. Raises `OfficeDocError` on the first problem."""
    layouts = manifest()['layouts']
    if not isinstance(slides, list) or not slides:
        raise OfficeDocError('The deck needs at least one slide: a list of {"layout": ..., fields...}.')
    if len(slides) + 1 > MAX_SLIDES:
        raise OfficeDocError(f'The deck has {len(slides)} slides plus the cover; the limit is {MAX_SLIDES} in all.')
    caps = capacities()
    out = []
    for n, slide in enumerate(slides, start=2):
        where = f'Slide {n}'
        if not isinstance(slide, dict):
            raise OfficeDocError(f'{where} must be an object with "layout" and that layout\'s fields.')
        key = str(slide.get('layout') or '').strip().lower()
        spec = layouts.get(key)
        if spec is None:
            raise OfficeDocError(f'{where} has layout "{slide.get("layout")}"; use one of: {", ".join(layouts)}.')
        where = f'Slide {n} ({key})'
        box = caps[spec['layout']]
        allowed = {'layout', 'notes', *spec['fields'], *spec.get('groups', {})}
        if spec.get('source'):
            allowed.add('source')
        if 'section_number' not in spec:
            allowed.add('section')
        for extra in ('table', 'chart'):
            if spec.get(extra):
                allowed.add(extra)
        unknown = sorted(set(slide) - allowed)
        if unknown:
            raise OfficeDocError(
                f'{where} has unknown keys: {", ".join(unknown)}. This layout takes: {", ".join(sorted(allowed - {"layout"}))}.'
            )

        clean: dict[str, Any] = {'layout': key, 'fill': {}, 'groups': {}}
        for name, field_spec in spec['fields'].items():
            value = slide.get(name)
            if value is None or (isinstance(value, (str, list)) and not value):
                if field_spec.get('required'):
                    raise OfficeDocError(f'{where} needs "{name}".')
                if field_spec.get('default'):
                    clean['fill'][field_spec['idx']] = field_spec['default']
                continue
            text = _field_text(
                value, f'{where} "{name}"', field_spec.get('list', False), field_spec.get('spaced', False)
            )
            _fit(text, box[field_spec['idx']], f'{where} "{name}"')
            clean['fill'][field_spec['idx']] = text
        if 'subtitle' in spec['fields'] and slide.get('subtitle') and 0 in clean['fill']:
            # A two-line title runs into the subtitle box directly under it.
            _fit(clean['fill'][0], Box(box[0].chars, 1), f'{where} "title" (with a subtitle, one line)')
        if spec.get('min_filled') and not any(name in slide and slide[name] for name in spec['min_filled']):
            raise OfficeDocError(f'{where} needs text for at least one of: {", ".join(spec["min_filled"])}.')

        for group, group_spec in spec.get('groups', {}).items():
            items = slide.get(group)
            low, high = group_spec['min'], group_spec['max']
            count = f'exactly {low}' if low == high else f'{low} to {high}'
            if not isinstance(items, list) or not low <= len(items) <= high:
                raise OfficeDocError(f'{where} needs "{group}": a list of {count} items.')
            slot_fields = [k for k in group_spec['slots'][0] if k != group_spec.get('number')]
            for i, (item, slot) in enumerate(zip(items, group_spec['slots']), start=1):
                if isinstance(item, str) and slot_fields == ['text']:
                    item = {'text': item}
                if not isinstance(item, dict):
                    raise OfficeDocError(
                        f'{where} "{group}" item {i} must be an object with: {", ".join(slot_fields)}.'
                    )
                unknown = sorted(set(item) - set(slot_fields))
                if unknown:
                    raise OfficeDocError(
                        f'{where} "{group}" item {i} has unknown keys: {", ".join(unknown)}; it takes {", ".join(slot_fields)}.'
                    )
                for name in slot_fields:
                    value = item.get(name)
                    if value is None or (isinstance(value, str) and not value.strip()):
                        raise OfficeDocError(f'{where} "{group}" item {i} needs "{name}".')
                    text = _field_text(
                        value if isinstance(value, str) else str(value), f'{where} "{group}" item {i} "{name}"', False
                    )
                    _fit(text, box[slot[name]], f'{where} "{group}" item {i} "{name}"')
                    clean['fill'][slot[name]] = text
                if group_spec.get('number'):
                    clean['fill'][slot[group_spec['number']]] = f'{i:02d}'
            clean['groups'][group] = len(items)

        if 'section_number' not in spec:
            section = slide.get('section')
            if section:
                text = _text(section, f'{where} "section"', 40).upper()
                _fit(text, box[2], f'{where} "section"')
                clean['fill'][2] = text
            else:
                clean['fill'][2] = spec.get('section') or section_labels()[spec['layout']]

        source_spec = spec.get('source')
        if source_spec:
            source = slide.get('source')
            if source:
                text = _text(source, f'{where} "source"', MAX_FIELD)
                _fit(text, box[source_spec['idx']], f'{where} "source"')
                clean['fill'][source_spec['idx']] = text
            elif source_spec.get('required'):
                raise OfficeDocError(
                    f'{where} needs "source": where the numbers come from and their period, for example '
                    '"Sumber: RKA 2027 (Direksi). Periode: 2027." The template requires it on data slides.'
                )

        if spec.get('table'):
            clean['table'] = _table(slide.get('table'), where, spec['table']['rows'])
        if spec.get('chart'):
            clean['chart'] = _chart(slide.get('chart'), where)
        notes = slide.get('notes')
        clean['notes'] = _text(notes, f'{where} "notes"', MAX_NOTES, required=False) if notes else ''
        out.append(clean)
    return out


# -- rendering ---------------------------------------------------------------------------------


def _set_text(shape_el, layout_el, text: str) -> None:
    """Write `text` into a slide placeholder with the layout prompt's paragraph and run formatting."""
    proto_p = next(iter(layout_el.iter(qn('a:p'))), None) if layout_el is not None else None
    proto_ppr = proto_p.find(qn('a:pPr')) if proto_p is not None else None
    proto_rpr = next(iter(proto_p.iter(qn('a:rPr'))), None) if proto_p is not None else None
    tx = shape_el.find(qn('p:txBody'))
    if tx is None:
        tx = etree.SubElement(shape_el, qn('p:txBody'))
        etree.SubElement(tx, qn('a:bodyPr'))
        etree.SubElement(tx, qn('a:lstStyle'))
    for p in tx.findall(qn('a:p')):
        tx.remove(p)
    for line in text.split('\n'):
        p = etree.SubElement(tx, qn('a:p'))
        if proto_ppr is not None:
            p.append(copy.deepcopy(proto_ppr))
        parts = _BOLD.split(line)
        for i, part in enumerate(parts):
            if not part:
                continue
            r = etree.SubElement(p, qn('a:r'))
            rpr = copy.deepcopy(proto_rpr) if proto_rpr is not None else etree.Element(qn('a:rPr'))
            if i % 2 == 1:
                rpr.set('b', '1')
            r.append(rpr)
            etree.SubElement(r, qn('a:t')).text = part
        if proto_rpr is not None:
            end = copy.deepcopy(proto_rpr)
            end.tag = qn('a:endParaRPr')
            p.append(end)


def _next_id(slide) -> str:
    """A shape id not yet used on `slide`: PowerPoint offers to repair a slide with two alike."""
    ids = [int(e.get('id')) for e in slide.shapes._spTree.iter(qn('p:cNvPr')) if e.get('id', '').isdigit()]
    return str(max(ids, default=1) + 1)


def _new_slide(prs, layout):
    """A slide on `layout`, with the slide-number placeholder python-pptx leaves out."""
    slide = prs.slides.add_slide(layout)
    present = set(_ph_by_idx(slide))
    for idx, el in _ph_by_idx(layout).items():
        ph = el.find(f'.//{{{_P}}}nvPr/{{{_P}}}ph')
        if idx not in present and ph.get('type') == 'sldNum':
            number = copy.deepcopy(el)
            number.find(f'.//{{{_P}}}cNvPr').set('id', _next_id(slide))
            slide.shapes._spTree.append(number)
    return slide


def _prune(slide, keep: set[int]) -> None:
    """Remove the slide placeholders nothing was written into, so no empty prompt is left behind."""
    tree = slide.shapes._spTree
    for idx, el in _ph_by_idx(slide).items():
        if idx not in keep:
            tree.remove(el)


def _cell_text(tc, text: str, align: str | None) -> None:
    tx = tc.find(qn('a:txBody'))
    proto_p = tx.find(qn('a:p'))
    proto_ppr = proto_p.find(qn('a:pPr')) if proto_p is not None else None
    proto_rpr = next(iter(proto_p.iter(qn('a:rPr'))), None) if proto_p is not None else None
    if proto_rpr is None and proto_p is not None:
        proto_rpr = proto_p.find(qn('a:endParaRPr'))
    for p in tx.findall(qn('a:p')):
        tx.remove(p)
    p = etree.SubElement(tx, qn('a:p'))
    ppr = copy.deepcopy(proto_ppr) if proto_ppr is not None else etree.Element(qn('a:pPr'))
    if align:
        ppr.set('algn', align)
    elif 'algn' in ppr.attrib:
        del ppr.attrib['algn']
    p.append(ppr)
    if text:
        r = etree.SubElement(p, qn('a:r'))
        rpr = copy.deepcopy(proto_rpr) if proto_rpr is not None else etree.Element(qn('a:rPr'))
        rpr.tag = qn('a:rPr')
        r.append(rpr)
        etree.SubElement(r, qn('a:t')).text = text
    if proto_rpr is not None:
        end = copy.deepcopy(proto_rpr)
        end.tag = qn('a:endParaRPr')
        p.append(end)


def _cell_box(width_emu: int, tc, size: float) -> Box:
    tcpr = tc.find(qn('a:tcPr'))
    margin = 0
    if tcpr is not None:
        margin = int(tcpr.get('marL', 91440)) + int(tcpr.get('marR', 91440))
    else:
        margin = 2 * 91440
    width_pt = max(1, width_emu - margin) / _EMU_PER_PT
    return Box(chars=max(1, math.floor(width_pt / (_EM['Space Grotesk'] * size))), lines=1)


def _font_size(tc, default: float) -> float:
    sizes = [int(e.get('sz')) / 100 for e in tc.iter(qn('a:rPr'), qn('a:endParaRPr'), qn('a:defRPr')) if e.get('sz')]
    return max(sizes) if sizes else default


def _put_table(slide, idx: int, proto_frame, table: dict, where: str) -> None:
    """Replace the table placeholder with the template's example table, rebuilt around `table`."""
    holder = _ph_by_idx(slide).get(idx)
    frame = copy.deepcopy(proto_frame)
    tbl = frame.find(f'.//{{{_A}}}tbl')
    proto_rows = tbl.findall(qn('a:tr'))
    head_proto, body_proto = proto_rows[0], proto_rows[-1]
    grid = tbl.find(qn('a:tblGrid'))
    proto_widths = [int(g.get('w')) for g in grid.findall(qn('a:gridCol'))]
    total = sum(proto_widths)
    n = len(table['columns'])
    if n == len(proto_widths):
        widths = proto_widths
    else:
        weights = [1.6] + [1.0] * (n - 1)
        widths = [int(total * w / sum(weights)) for w in weights]
        widths[-1] = total - sum(widths[:-1])
    for g in grid.findall(qn('a:gridCol')):
        grid.remove(g)
    for w in widths:
        etree.SubElement(grid, qn('a:gridCol'), w=str(w))
    for tr in proto_rows:
        tbl.remove(tr)

    # Rows keep the example's height: the layout draws a rule under each example row.
    row_height = int(head_proto.get('h'))
    aligns = ['r' if c['type'] in _NUMERIC else None for c in table['columns']]

    def build(proto_tr, texts: list[str], header: bool) -> None:
        tr = copy.deepcopy(proto_tr)
        tr.set('h', str(row_height))
        cells = tr.findall(qn('a:tc'))
        for tc in cells:
            tr.remove(tc)
        for c, text in enumerate(texts):
            tc = copy.deepcopy(cells[0] if c == 0 else cells[min(1, len(cells) - 1)])
            _cell_text(tc, text, aligns[c])
            if text:
                size = _font_size(tc, 18.0)
                box = _cell_box(widths[c], tc, size)
                label = f'{where} table header "{text}"' if header else f'{where} table cell "{text[:40]}"'
                _fit(text, box, label)
            tr.append(tc)
        tbl.append(tr)

    build(head_proto, [c['header'].upper() for c in table['columns']], True)
    for row in table['rows']:
        build(body_proto, row, False)

    frame.find(f'.//{{{_P}}}cNvPr').set('id', _next_id(slide))
    tree = slide.shapes._spTree
    if holder is not None:
        holder.addprevious(frame)
        tree.remove(holder)
    else:
        tree.append(frame)


def _put_chart(slide, idx: int, proto_space, kind: str, chart: dict) -> None:
    """A native chart in the chart placeholder, styled as the template's example chart."""
    holder = next(s for s in slide.placeholders if s.placeholder_format.idx == idx)
    whole = all(
        isinstance(v, int) or (isinstance(v, float) and v.is_integer())
        for s in chart['series']
        for v in s['values']
        if v is not None
    )
    data = CategoryChartData(number_format='#,##0' if whole else '#,##0.0')
    data.categories = chart['categories']
    for series in chart['series']:
        data.add_series(series['name'], series['values'])
    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if kind == 'bar' else XL_CHART_TYPE.LINE_MARKERS
    frame = holder.insert_chart(chart_type, data)
    part = frame.chart_part
    xlsx_rid = next(rid for rid, rel in part.rels.items() if rel.reltype == RT.PACKAGE)
    styled = copy.deepcopy(proto_space)
    external = styled.find(qn('c:externalData'))
    if external is not None:
        external.set(qn('r:id'), xlsx_rid)
    for stale in ('c:userShapes',):
        for el in styled.findall(qn(stale)):
            styled.remove(el)
    # The example fixes its value axis to its own data (0-150, 0-120); let it follow this data.
    values = [v for s in chart['series'] for v in s['values'] if v is not None]
    for scaling in styled.iter(qn('c:scaling')):
        for bound in ('c:max', 'c:min'):
            el = scaling.find(qn(bound))
            if el is not None and (bound == 'c:max' or min(values, default=0) < 0):
                scaling.remove(el)
    for unit in list(styled.iter(qn('c:majorUnit'))):
        unit.getparent().remove(unit)
    part._element = styled
    part.__dict__.pop('chart', None)
    part.__dict__.pop('chart_workbook', None)
    native = part.chart
    proto_count = len(native.plots[0].series)
    native.replace_data(data)
    for i, series in enumerate(native.plots[0].series):
        if i < proto_count:
            continue
        color = RGBColor.from_string(_PALETTE[i % len(_PALETTE)])
        if kind == 'bar':
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        else:
            series.format.line.color.rgb = color
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
    if len(chart['series']) > 1 and not native.has_legend:
        native.has_legend = True
        native.legend.position = XL_LEGEND_POSITION.BOTTOM
        native.legend.include_in_layout = False
        native.legend.font.name = 'Space Grotesk'
        native.legend.font.size = 15 * _EMU_PER_PT
    if len(chart['series']) == 1 and native.has_legend:
        native.has_legend = False
    # The example chart's alt text calls the data illustrative; this one is not.
    frame._element.find(f'.//{{{_P}}}cNvPr').set(
        'descr', 'Grafik yang dapat diedit. Data dan sumbernya tercantum pada slide.'
    )


def _prototypes(prs) -> tuple[dict[str, Any], dict[str, Any]]:
    """The example slides' table frames and chart XML, by layout name."""
    tables, charts = {}, {}
    for slide in prs.slides:
        name = slide.slide_layout.name
        for el in slide.shapes._spTree.iterchildren(qn('p:graphicFrame')):
            uri = el.find(f'.//{{{_A}}}graphicData').get('uri')
            if uri.endswith('/table'):
                tables.setdefault(name, copy.deepcopy(el))
            elif uri.endswith('/chart'):
                rid = el.find(f'.//{{{_A}}}graphicData/').get(qn('r:id'))
                charts.setdefault(name, copy.deepcopy(slide.part.related_part(rid)._element))
    return tables, charts


def _drop_example_slides(prs) -> None:
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        rid = sld.get(qn('r:id'))
        ids.remove(sld)
        prs.part.drop_rel(rid)


def _footer_text(level: str, unit: str) -> str:
    return f'{unit} / {level} · {STATUS}' if unit else f'{level} · {STATUS}'


def _fill(slide, layout, fills: dict[int, str], fixed: list[int]) -> set[int]:
    layout_ph = _ph_by_idx(layout)
    slide_ph = _ph_by_idx(slide)
    written = set()
    for idx in fixed + [WEBSITE_IDX]:
        if idx in slide_ph and idx not in fills:
            fills = {**fills, idx: _layout_text(layout, idx)}
    for idx, text in fills.items():
        if idx in slide_ph:
            _set_text(slide_ph[idx], layout_ph.get(idx), text)
            written.add(idx)
    return written


def filename_for(title: str) -> str:
    slug = re.sub(r'[^A-Za-z0-9]+', '-', title).strip('-')[:60] or 'presentasi'
    return f'JUP-PPT-{slug}-Draf.pptx'


def render_slides(
    title: Any,
    slides: Any,
    *,
    cover: Any = None,
    classification: Any = None,
    floor: Any = None,
    author: str = '',
    now: dt.datetime | None = None,
) -> SlidesRendered:
    """Build a deck from a spec. Raises `OfficeDocError` for anything it will not render."""
    clean_title = _text(title, 'The title', MAX_TITLE)
    level = classification_for(classification, floor)
    front = _cover(cover, clean_title)
    clean = validate_slides(slides)
    caps = capacities()
    cover_spec = manifest()['covers'][front['style']]
    _fit(clean_title, caps[cover_spec['layout']][cover_spec['fields']['title']], 'The title on the cover')
    moment = (now or dt.datetime.now(JAKARTA)).astimezone(JAKARTA)

    prs = _open_template()
    layouts = _layouts_by_name(prs)
    tables, charts = _prototypes(prs)
    _drop_example_slides(prs)
    footer = _footer_text(level, front['unit'])

    # The cover.
    layout = layouts[cover_spec['layout']]
    slide = _new_slide(prs, layout)
    unit_date = '   '.join(v for v in (front['unit'], front['date'] or moment.strftime('%d/%m/%Y')) if v)
    fills = {
        cover_spec['fields']['title']: clean_title,
        cover_spec['fields']['unit_date']: unit_date,
        cover_spec['footer']: footer,
    }
    if front['kind']:
        _fit(front['kind'], caps[cover_spec['layout']][cover_spec['fields']['kind']], 'The cover "kind"')
        fills[cover_spec['fields']['kind']] = front['kind']
    _prune(slide, _fill(slide, layout, fills, cover_spec.get('fixed', [])))

    unfilled: list[str] = []
    sections = 0
    chart_count = table_count = 0
    for n, spec_slide in enumerate(clean, start=2):
        spec = manifest()['layouts'][spec_slide['layout']]
        layout = layouts[spec['layout']]
        slide = _new_slide(prs, layout)
        fills = dict(spec_slide['fill'])
        if 'section_number' in spec:
            sections += 1
            fills[spec['section_number']] = f'{sections:02d}'
        if FOOTER_IDX in caps[spec['layout']]:
            fills[FOOTER_IDX] = footer
        keep = _fill(slide, layout, fills, spec.get('fixed', []))
        keep |= {idx for idx, el in _ph_by_idx(slide).items() if el.find(f'.//{{{_P}}}ph').get('type') == 'sldNum'}
        if spec.get('photo'):
            keep.add(spec['photo']['idx'])
            unfilled.append(f'slide {n}: the project photo')
        if spec.get('table'):
            keep.add(spec['table']['idx'])
        if spec.get('chart'):
            keep.add(spec['chart']['idx'])
        _prune(slide, keep)
        if spec.get('table'):
            _put_table(
                slide,
                spec['table']['idx'],
                tables[spec['layout']],
                spec_slide['table'],
                f'Slide {n} ({spec_slide["layout"]})',
            )
            table_count += 1
        if spec.get('chart'):
            _put_chart(slide, spec['chart']['idx'], charts[spec['layout']], spec['chart']['kind'], spec_slide['chart'])
            chart_count += 1
        if spec_slide['notes']:
            slide.notes_slide.notes_text_frame.text = spec_slide['notes']

    props = prs.core_properties
    props.title = clean_title
    props.author = author[:200] or 'Asisten JUP'
    props.last_modified_by = 'Asisten JUP'
    props.keywords = 'JUP-PPT'
    props.subject = f'{STATUS} · {level}'
    props.comments = f'Draf dibuat oleh Asisten JUP. Klasifikasi: {level}.'
    props.created = props.modified = moment.replace(tzinfo=None)
    props.revision = 1

    out = io.BytesIO()
    prs.save(out)
    data = out.getvalue()
    if len(data) > MAX_BYTES:
        raise OfficeDocError(f'The deck would be {len(data)} bytes; the limit is {MAX_BYTES}. Use fewer slides.')
    return SlidesRendered(
        data=data,
        filename=filename_for(clean_title),
        title=clean_title,
        classification=level,
        slides=len(clean) + 1,
        charts=chart_count,
        tables=table_count,
        unfilled=unfilled,
    )


# -- the model's guide -------------------------------------------------------------------------


def _describe_field(name: str, box: Box, extra: str = '') -> str:
    return f'{name} ({box.describe()}{extra})'


def layout_guide(key: Any = '') -> str:
    """The layout list, or one layout's fields with what each box holds."""
    layouts = manifest()['layouts']
    caps = capacities()
    wanted = str(key or '').strip().lower()
    if not wanted:
        lines = [
            'JUP slide layouts (Template presentasi JUP v1.5-r2). The deck always starts with a cover the server '
            'builds from the title and "cover" ({"style": "institusional" | "korporat", "kind", "unit", "date"}). '
            'Ask for one layout here to see its fields and how much text each holds.'
        ]
        for name, spec in layouts.items():
            lines.append(f'- {name}: {spec["name"]}')
        return '\n'.join(lines)
    spec = layouts.get(wanted)
    if spec is None:
        return f'Unknown layout "{key}". Use one of: {", ".join(layouts)}.'
    box = caps[spec['layout']]
    parts = [f'{wanted}: {spec["name"]} (template layout "{spec["layout"]}").', 'Fields:']
    for name, field_spec in spec['fields'].items():
        flags = []
        if field_spec.get('required'):
            flags.append('required')
        if field_spec.get('default'):
            flags.append(f'default "{field_spec["default"]}"')
        if field_spec.get('list'):
            flags.append('a string or a list of lines')
        if field_spec.get('spaced'):
            flags.append('a blank line is put between items and counts as a line')
        if name == 'title' and 'subtitle' in spec['fields']:
            flags.append('one line only when a subtitle is given')
        parts.append('- ' + _describe_field(name, box[field_spec['idx']], ('; ' + ', '.join(flags)) if flags else ''))
    if spec.get('min_filled'):
        parts.append(
            f'At least one of {", ".join(spec["min_filled"])} must be given; the sector names and icons are fixed.'
        )
    for group, group_spec in spec.get('groups', {}).items():
        slot = group_spec['slots'][0]
        names = [k for k in slot if k != group_spec.get('number')]
        count = (
            f'exactly {group_spec["min"]}'
            if group_spec['min'] == group_spec['max']
            else f'{group_spec["min"]} to {group_spec["max"]}'
        )
        inner = ', '.join(_describe_field(k, box[slot[k]]) for k in names)
        numbered = '; numbered 01, 02, … automatically' if group_spec.get('number') else ''
        shape = '"text" strings or objects' if names == ['text'] else 'objects'
        parts.append(f'- {group}: {count} {shape} with {inner}{numbered}')
    if 'section_number' in spec:
        parts.append('- The section number (01, 02, …) is counted automatically.')
    else:
        default = spec.get('section') or ''
        parts.append(
            f'- section (optional label above the title, {box[2].describe()}{f"; default {default!r}" if default else ""})'
        )
    if spec.get('table'):
        parts.append(
            f'- table: {{"columns": [{{"header", "type"}}], "rows": [[...]]}}, 2 to {MAX_TABLE_COLS} columns and 1 to '
            f'{spec["table"]["rows"]} rows; types {", ".join(TABLE_TYPES)} (numbers as JSON numbers, dates YYYY-MM-DD); '
            'each cell is one line, so keep cells short and state the unit in the subtitle.'
        )
    if spec.get('chart'):
        kind = 'bar (clustered columns)' if spec['chart']['kind'] == 'bar' else 'line with markers'
        parts.append(
            f'- chart ({kind}, editable in PowerPoint): {{"categories": [...], "series": [{{"name", "values"}}]}}, '
            f'2 to {MAX_CATEGORIES} categories and 1 to {MAX_SERIES} series of numbers (null for a gap); state the unit in the subtitle.'
        )
    if spec.get('source'):
        required = 'required' if spec['source'].get('required') else 'optional'
        parts.append(
            f'- source ({required}, {box[spec["source"]["idx"]].describe()}): where the content comes from and its period.'
        )
    if spec.get('photo'):
        parts.append('- The photo frame is left for the user to fill.')
    parts.append(f'- notes (optional speaker notes, up to {MAX_NOTES} characters).')
    return '\n'.join(parts)


__all__ = [
    'MAX_SLIDES',
    'PPTX_MIME',
    'SlidesRendered',
    'capacities',
    'filename_for',
    'layout_guide',
    'lines_needed',
    'manifest',
    'render_slides',
    'template_path',
    'validate_slides',
]
