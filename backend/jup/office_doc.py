"""Word documents on the JUP templates (Track O, O1 1/3, #44).

The assistant never writes a file itself. It sends a **spec** (a template code, a title, optional
metadata, a classification and a list of blocks), and this module fills the chosen template from
JUP Template Dokumen 2.0 with it. There is no model-written code anywhere, so nothing here needs
a sandbox, and nothing leaves the server.

**What is kept from the template:** its styles, its header and footer, and its head block (type
label, title, and the metadata lines for number, version, owner, date, distribution and status).
The template's example body is replaced by the spec's blocks, written in the template's own
styles (`Heading1-3`, `ListBullet`, `ListNumber`, `JUPTable`, `JUPNote`). `outline()` gives the
model the template's intended structure so its blocks can follow it.

**Rules the model cannot override** (each tested in `tests/test_jup_office_doc.py`):

- The status is always **Draf**: a model has no proof of approval, which the template requires.
- The classification is INTERNAL, TERBATAS or RAHASIA, INTERNAL by default. **PUBLIK is
  refused**: publication needs an approval a draft does not have. A caller-supplied `floor`
  raises the classification and never lowers it; that is the hook the contract and lawsuit
  tools use, so a file built from restricted material is at least as restricted as its source.
- Metadata that is not given stays as the template's `[bracketed]` prompt. Nothing is invented.
- Limits are refused with a message, never silently truncated.

Pure: no app imports, no network, no disk writes. The chat tool (`jup.office_tool`) stores the
bytes this returns.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from lxml import etree

TEMPLATE_DIR = Path(__file__).resolve().parent / 'office_templates'
DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

STATUS = 'Draf'
# Lowest first. PUBLIK is a real level in the template, but never one a draft may carry.
LEVELS = ('INTERNAL', 'TERBATAS', 'RAHASIA')
PUBLIC = 'PUBLIK'
DEFAULT_LEVEL = 'INTERNAL'

MAX_BLOCKS = 200
MAX_CHARS = 60_000
MAX_TEXT = 4_000
MAX_ITEMS = 100
MAX_TABLE_COLS = 12
MAX_TABLE_ROWS = 200
MAX_TITLE = 200
MAX_META = 200
MAX_BYTES = 5 * 1024 * 1024

BLOCK_TYPES = ('heading', 'paragraph', 'bullets', 'numbered', 'table', 'note')
HEAD_STYLES = ('JUPEyebrow', 'Title', 'JUPMetadata')
# The metadata a spec may fill, by the position of the control in the template's head block.
META_FIELDS = ('unit', 'number', 'version', 'owner', 'date', 'distribution')
_HEAD_LAYOUT = (('unit', 'number', 'version'), ('owner', 'date'), ('distribution',))

_BOLD = re.compile(r'\*\*(.+?)\*\*')
_NS = {'w': nsmap['w']}


def _xp(el, expr: str):
    """XPath with the `w:` prefix bound, for plain lxml elements as well as python-docx ones."""
    return etree.XPath(expr, namespaces=_NS)(el)


class OfficeDocError(ValueError):
    """A spec this module will not render. The message is written for the model to relay."""


@dataclass(frozen=True)
class Template:
    code: str
    file: str
    name: str
    group: str
    sha256: str

    @property
    def path(self) -> Path:
        return TEMPLATE_DIR / 'docx' / self.file


@dataclass
class Rendered:
    data: bytes
    filename: str
    template: Template
    classification: str
    status: str = STATUS
    blocks: int = 0
    chars: int = 0
    unfilled: list[str] = field(default_factory=list)


@lru_cache(maxsize=1)
def templates() -> dict[str, Template]:
    """The templates by code, from `office_templates/manifest.json`."""
    manifest = json.loads((TEMPLATE_DIR / 'manifest.json').read_text(encoding='utf-8'))
    return {t['code']: Template(**t) for t in manifest['templates']}


def template_for(code: Any) -> Template:
    key = str(code or '').strip().upper()
    found = templates().get(key)
    if found is None:
        known = ', '.join(f'{t.code} ({t.name})' for t in templates().values())
        raise OfficeDocError(f'Unknown document template "{code}". Use one of: {known}.')
    return found


def classification_for(requested: Any, floor: Any = None) -> str:
    """The classification a draft carries: the requested one, raised to `floor`, never PUBLIK."""
    level = str(requested or DEFAULT_LEVEL).strip().upper()
    if level == PUBLIC:
        raise OfficeDocError(
            'A generated document cannot be classified PUBLIK: publication needs an approval that a '
            'draft does not have. Use INTERNAL (the default), TERBATAS or RAHASIA.'
        )
    if level not in LEVELS:
        raise OfficeDocError(f'Unknown classification "{requested}". Use INTERNAL, TERBATAS or RAHASIA.')
    if floor is not None:
        floor_level = str(floor).strip().upper()
        if floor_level not in LEVELS:
            raise OfficeDocError(f'Unknown classification floor "{floor}".')
        if LEVELS.index(floor_level) > LEVELS.index(level):
            level = floor_level
    return level


# -- spec validation ---------------------------------------------------------------------------


def _text(value: Any, what: str, limit: int = MAX_TEXT, required: bool = True) -> str:
    if value is None:
        value = ''
    if not isinstance(value, str):
        raise OfficeDocError(f'{what} must be text.')
    value = value.strip()
    if required and not value:
        raise OfficeDocError(f'{what} is empty.')
    if len(value) > limit:
        raise OfficeDocError(f'{what} is {len(value)} characters; the limit is {limit}. Split it or shorten it.')
    return value


def _items(value: Any, what: str) -> list[str]:
    if not isinstance(value, list) or not value:
        raise OfficeDocError(f'{what} needs a non-empty list of items.')
    if len(value) > MAX_ITEMS:
        raise OfficeDocError(f'{what} has {len(value)} items; the limit is {MAX_ITEMS}.')
    return [_text(item, f'An item in {what}') for item in value]


def validate_blocks(blocks: Any) -> list[dict]:
    """The blocks, checked and normalised. Raises `OfficeDocError` on the first problem found."""
    if not isinstance(blocks, list) or not blocks:
        raise OfficeDocError('The document has no content: send at least one block.')
    if len(blocks) > MAX_BLOCKS:
        raise OfficeDocError(f'The document has {len(blocks)} blocks; the limit is {MAX_BLOCKS}.')
    out: list[dict] = []
    for n, block in enumerate(blocks, start=1):
        where = f'Block {n}'
        if not isinstance(block, dict):
            raise OfficeDocError(f'{where} must be an object with a "type".')
        kind = block.get('type')
        if kind not in BLOCK_TYPES:
            raise OfficeDocError(f'{where} has type "{kind}"; use one of: {", ".join(BLOCK_TYPES)}.')
        if kind == 'heading':
            level = block.get('level', 2)
            if level not in (1, 2, 3):
                raise OfficeDocError(f'{where}: a heading level is 1, 2 or 3.')
            out.append({'type': kind, 'level': level, 'text': _text(block.get('text'), f'{where} heading', MAX_TITLE)})
        elif kind in ('paragraph', 'note'):
            out.append({'type': kind, 'text': _text(block.get('text'), f'{where} text')})
        elif kind in ('bullets', 'numbered'):
            out.append({'type': kind, 'items': _items(block.get('items'), f'{where} list')})
        else:
            header = block.get('header')
            rows = block.get('rows')
            if not isinstance(header, list) or not header:
                raise OfficeDocError(f'{where}: a table needs a non-empty "header" list.')
            if len(header) > MAX_TABLE_COLS:
                raise OfficeDocError(f'{where}: a table has at most {MAX_TABLE_COLS} columns.')
            if not isinstance(rows, list) or not rows:
                raise OfficeDocError(f'{where}: a table needs at least one row in "rows".')
            if len(rows) > MAX_TABLE_ROWS:
                raise OfficeDocError(f'{where}: a table has at most {MAX_TABLE_ROWS} rows.')
            cols = len(header)
            clean_rows = []
            for r, row in enumerate(rows, start=1):
                if not isinstance(row, list) or len(row) != cols:
                    raise OfficeDocError(f'{where}, row {r}: every row needs exactly {cols} cells, like the header.')
                clean_rows.append([_text(c, f'{where}, row {r} cell', required=False) for c in row])
            out.append(
                {
                    'type': kind,
                    'header': [_text(h, f'{where} header cell') for h in header],
                    'rows': clean_rows,
                }
            )
    total = sum(_block_chars(b) for b in out)
    if total > MAX_CHARS:
        raise OfficeDocError(f'The document has {total} characters of content; the limit is {MAX_CHARS}.')
    return out


def _block_chars(block: dict) -> int:
    if 'text' in block:
        return len(block['text'])
    if 'items' in block:
        return sum(len(i) for i in block['items'])
    return sum(len(h) for h in block['header']) + sum(len(c) for row in block['rows'] for c in row)


# -- template filling --------------------------------------------------------------------------


def _style_id(el) -> str:
    return _xp(el, 'string(./w:pPr/w:pStyle/@w:val)')


def _head(document: DocumentObject) -> list:
    """The leading paragraphs in a head style: type label, title, metadata lines."""
    head = []
    for el in document.element.body.iterchildren():
        if el.tag != qn('w:p') or _style_id(el) not in HEAD_STYLES:
            break
        head.append(el)
    return head


def _set_sdt(sdt, text: str) -> None:
    """Replace a content control's shown text, keeping the first run's formatting."""
    content = sdt.find(qn('w:sdtContent'))
    runs = content.findall('.//' + qn('w:r'))
    first = next((r for r in runs if r.find(qn('w:t')) is not None), None)
    if first is None:
        raise OfficeDocError('The template has a content control without text; the template may be damaged.')
    for run in runs:
        if run is not first:
            run.getparent().remove(run)
    for t in first.findall(qn('w:t'))[1:]:
        first.remove(t)
    t = first.find(qn('w:t'))
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    rpr = first.find(qn('w:rPr'))
    if rpr is not None:
        for rstyle in rpr.findall(qn('w:rStyle')):
            if rstyle.get(qn('w:val')) == 'PlaceholderText':
                rpr.remove(rstyle)
    plc = sdt.find(qn('w:sdtPr') + '/' + qn('w:showingPlcHdr'))
    if plc is not None:
        plc.getparent().remove(plc)


def _sdts(el) -> list:
    return el.findall('.//' + qn('w:sdt'))


def _tagged(root, tag: str) -> list:
    return [s for s in _sdts(root) if _xp(s, 'string(./w:sdtPr/w:tag/@w:val)') == tag]


def _footers(document: DocumentObject) -> list:
    return [rel.target_part.element for rel in document.part.rels.values() if rel.reltype == RT.FOOTER]


def _fill_head(document: DocumentObject, title: str, meta: dict[str, str]) -> list[str]:
    """Fill the head block. Returns the names of metadata left as the template's prompt."""
    head = _head(document)
    titles = [el for el in head if _style_id(el) == 'Title']
    metas = [el for el in head if _style_id(el) == 'JUPMetadata']
    if len(titles) != 1 or len(_sdts(titles[0])) != 1 or len(metas) < 4:
        raise OfficeDocError('The template does not have the JUP head block; it cannot be filled.')
    _set_sdt(_sdts(titles[0])[0], title)
    unfilled = []
    for line, names in zip(metas, _HEAD_LAYOUT):
        controls = _sdts(line)
        if len(controls) != len(names):
            raise OfficeDocError('The template head block has an unexpected shape; it cannot be filled.')
        for control, name in zip(controls, names):
            if meta.get(name):
                _set_sdt(control, meta[name])
            else:
                unfilled.append(name)
    statuses = _tagged(document.element.body, 'JUP-STATUS')
    if not statuses:
        raise OfficeDocError('The template has no status field; it cannot be marked as a draft.')
    for control in statuses:
        _set_sdt(control, STATUS)
    return unfilled


def _fill_footer(document: DocumentObject, classification: str, group: str) -> None:
    access = [c for footer in _footers(document) for c in _tagged(footer, 'JUP-ACCESS')]
    if not access:
        raise OfficeDocError('The template footer has no classification field.')
    for control in access:
        _set_sdt(control, classification)
    for footer in _footers(document):
        for control in _tagged(footer, 'JUP-GROUP'):
            _set_sdt(control, group)


def _clear_body(document: DocumentObject) -> None:
    """Remove everything after the head block except the final section properties."""
    body = document.element.body
    # The list keeps the head's proxies alive: lxml then hands back the same objects, so `in`
    # compares identity. An `id()` of a discarded proxy can be reused by an unrelated element.
    head = _head(document)
    for el in list(body.iterchildren()):
        if el in head or el.tag == qn('w:sectPr'):
            continue
        body.remove(el)


def _style_name(document: DocumentObject, style_id: str) -> str:
    for style in document.styles:
        if style.style_id == style_id:
            return style.name
    raise OfficeDocError(f'The template is missing the "{style_id}" style.')


def _add_runs(paragraph, text: str) -> None:
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _write_blocks(document: DocumentObject, blocks: list[dict]) -> None:
    names = {
        sid: _style_name(document, sid)
        for sid in ('Heading1', 'Heading2', 'Heading3', 'ListBullet', 'ListNumber', 'JUPNote', 'JUPTable')
    }
    for block in blocks:
        kind = block['type']
        if kind == 'heading':
            document.add_paragraph(block['text'], style=names[f'Heading{block["level"]}'])
        elif kind == 'paragraph':
            _add_runs(document.add_paragraph(), block['text'])
        elif kind == 'note':
            _add_runs(document.add_paragraph(style=names['JUPNote']), block['text'])
        elif kind in ('bullets', 'numbered'):
            style = names['ListBullet' if kind == 'bullets' else 'ListNumber']
            for item in block['items']:
                _add_runs(document.add_paragraph(style=style), item)
        else:
            _add_table(document, names['JUPTable'], block)


# The template's tables carry direct formatting rather than a table style, so a new table copies
# it: header row filled 004999 with white bold text, rows alternating FFFFFF and F0F1F1 with 011627
# text, 10.5 pt, D9D9D9 borders, the template's cell padding, the header repeated on each page, and `JUPTable` paragraphs.
_HEADER_FILL, _ROW_FILLS, _BORDER = '004999', ('FFFFFF', 'F0F1F1'), 'D9D9D9'
_HEADER_TEXT, _BODY_TEXT, _TABLE_SIZE = 'FFFFFF', '011627', '21'


def _el(tag: str, **attrs: str):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f'w:{key}'), value)
    return el


def _content_width(document: DocumentObject) -> int:
    section = document.sections[-1]
    return int((section.page_width - section.left_margin - section.right_margin) / 635)  # EMU -> twips


def _format_cell(cell, text: str, fill: str, color: str, bold: bool, style: str, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:tcW')) + tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    tc_pr.append(_el('w:tcW', type='dxa', w=str(width)))
    tc_pr.append(_el('w:vAlign', val='center'))
    margins = _el('w:tcMar')
    for side, size in (('top', '95'), ('left', '105'), ('bottom', '95'), ('right', '105')):
        margins.append(_el(f'w:{side}', w=size, type='dxa'))
    tc_pr.append(margins)
    tc_pr.append(_el('w:shd', val='clear', color='auto', fill=fill))
    paragraph = cell.paragraphs[0]
    paragraph.style = style
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    pos = 0
    pieces = []
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            pieces.append((text[pos : match.start()], bold))
        pieces.append((match.group(1), True))
        pos = match.end()
    if pos < len(text) or not pieces:
        pieces.append((text[pos:], bold))
    for piece, strong in pieces:
        run = paragraph.add_run(piece)
        run.bold = strong
        r_pr = run._r.get_or_add_rPr()
        r_pr.append(_el('w:color', val=color))
        r_pr.append(_el('w:sz', val=_TABLE_SIZE))


def _add_table(document: DocumentObject, style: str, block: dict) -> None:
    cols = len(block['header'])
    table = document.add_table(rows=1 + len(block['rows']), cols=cols)
    width = _content_width(document) // cols
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(_el('w:jc', val='center'))
    borders = _el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(_el(f'w:{side}', val='single', sz='4', color=_BORDER))
    tbl_pr.append(borders)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(_el('w:tblHeader', val='true'))
    for cell, text in zip(header.cells, block['header']):
        _format_cell(cell, text, _HEADER_FILL, _HEADER_TEXT, True, style, width)
    for n, (row, values) in enumerate(zip(table.rows[1:], block['rows'])):
        for cell, text in zip(row.cells, values):
            _format_cell(cell, text, _ROW_FILLS[n % 2], _BODY_TEXT, False, style, width)
    document.add_paragraph()


def filename_for(template: Template, title: str) -> str:
    slug = re.sub(r'[^A-Za-z0-9]+', '-', title).strip('-')[:60] or 'dokumen'
    return f'JUP-{template.code}-{slug}-Draf.docx'


def render(
    template: Any,
    title: Any,
    blocks: Any,
    *,
    classification: Any = None,
    floor: Any = None,
    metadata: dict | None = None,
    author: str = '',
) -> Rendered:
    """Fill a JUP template with a spec. Raises `OfficeDocError` for anything it will not render."""
    tpl = template_for(template)
    clean_title = _text(title, 'The title', MAX_TITLE)
    level = classification_for(classification, floor)
    clean_blocks = validate_blocks(blocks)
    meta = {}
    for name in META_FIELDS:
        value = (metadata or {}).get(name)
        if value not in (None, ''):
            meta[name] = _text(value, f'The {name}', MAX_META)
    unknown = sorted(set(metadata or {}) - set(META_FIELDS))
    if unknown:
        raise OfficeDocError(f'Unknown metadata: {", ".join(unknown)}. Allowed: {", ".join(META_FIELDS)}.')

    raw = tpl.path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != tpl.sha256:
        raise OfficeDocError(f'The {tpl.name} template does not match its pinned copy; refusing to use it.')
    document = Document(io.BytesIO(raw))
    unfilled = _fill_head(document, clean_title, meta)
    _fill_footer(document, level, tpl.group)
    _clear_body(document)
    _write_blocks(document, clean_blocks)

    props = document.core_properties
    props.title = clean_title
    props.subject = tpl.name
    props.author = author[:MAX_META]
    props.last_modified_by = 'Asisten JUP'
    props.comments = f'Draf dibuat oleh Asisten JUP. Klasifikasi: {level}.'
    props.keywords = f'JUP-{tpl.code}'

    out = io.BytesIO()
    document.save(out)
    data = out.getvalue()
    if len(data) > MAX_BYTES:
        raise OfficeDocError(f'The document would be {len(data)} bytes; the limit is {MAX_BYTES}.')
    return Rendered(
        data=data,
        filename=filename_for(tpl, clean_title),
        template=tpl,
        classification=level,
        blocks=len(clean_blocks),
        chars=sum(_block_chars(b) for b in clean_blocks),
        unfilled=unfilled,
    )


# -- outline -----------------------------------------------------------------------------------


@lru_cache(maxsize=32)
def outline(code: Any) -> str:
    """A template's intended structure, for the model: its headings and the prompts under them."""
    tpl = template_for(code)
    document = Document(str(tpl.path))
    body = document.element.body
    head = _head(document)
    lines = [f'# {tpl.code}: {tpl.name}', f'Group: {tpl.group}', '']
    for el in body.iterchildren():
        if el in head:
            continue
        if el.tag == qn('w:p'):
            text = ''.join(_xp(el, './/w:t/text()')).strip()
            if not text:
                continue
            style = _style_id(el)
            if style == 'Heading1':
                lines.append(f'## {text}')
            elif style == 'Heading2':
                lines.append(f'### {text}')
            elif style == 'Heading3':
                lines.append(f'#### {text}')
            elif style in ('ListBullet', 'ListNumber'):
                lines.append(f'- {text}')
            else:
                lines.append(text)
        elif el.tag == qn('w:tbl'):
            first = el.find(qn('w:tr'))
            cells = (
                [''.join(_xp(c, './/w:t/text()')).strip() for c in first.findall(qn('w:tc'))]
                if first is not None
                else []
            )
            lines.append(f'[table: {" | ".join(cells)}]')
    return '\n'.join(lines).strip() + '\n'


def catalogue() -> str:
    """Every template, one line each, for the model to choose from."""
    return '\n'.join(f'- {t.code}: {t.name} ({t.group})' for t in templates().values())


__all__ = [
    'DOCX_MIME',
    'LEVELS',
    'OfficeDocError',
    'Rendered',
    'STATUS',
    'Template',
    'catalogue',
    'classification_for',
    'filename_for',
    'outline',
    'render',
    'template_for',
    'templates',
    'validate_blocks',
]
