"""The office chat tools (O1 2/3, #45): `jup.office_tool` and the stored wrapper `tools/jup_documents.py`.

Storage is the app's own upload path, which needs the app; these tests pass a fake `store` and
check everything around it. The live check (#46) proves the real upload, owner and 404.
"""

import ast
import asyncio
import io
from pathlib import Path

import pytest

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from jup import office_floor, office_tool
from jup.office_tool import (
    DECK_NOT_CREATED,
    FLOOR_FAILED,
    NO_CHAT,
    NO_USER,
    NOT_CREATED,
    SHEET_NOT_CREATED,
    STORE_FAILED,
    create_document,
    create_presentation,
    create_spreadsheet,
    document_templates,
    presentation_layouts,
)

TOOL_SOURCE = Path(__file__).resolve().parent.parent / 'tools' / 'jup_documents.py'
USER = {'id': 'u-1', 'name': 'Staf JUP', 'role': 'user'}
BLOCKS = [{'type': 'heading', 'level': 2, 'text': 'Agenda'}, {'type': 'bullets', 'items': ['Status SPBG']}]
SHEETS = [{'name': 'Kontrak', 'columns': [{'header': 'Mitra'}, {'header': 'Nilai', 'type': 'idr'}], 'rows': [['A', 5]]}]
SLIDES = [{'layout': 'agenda', 'title': 'Agenda', 'items': ['Kinerja', 'RKA 2027']}]


def floor_of(level):
    async def lookup(chat_id):
        lookup.asked.append(chat_id)
        return level

    lookup.asked = []
    return lookup


async def broken_floor(chat_id):
    raise RuntimeError('database unavailable')


class Recorder:
    def __init__(self, fail: bool = False):
        self.fail = fail
        self.stored = []
        self.events = []

    async def store(self, request, rendered, user, chat_id, message_id):
        if self.fail:
            raise OSError('bucket unavailable')
        self.stored.append((request, rendered, user, chat_id, message_id))
        return office_tool.file_item('file-1', rendered)

    async def emit(self, event):
        self.events.append(event)


def _run(recorder: Recorder, **overrides) -> str:
    kwargs = {
        'request': 'REQ',
        'user': USER,
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'template': 'NOT',
        'title': 'Rapat mingguan operasi',
        'blocks': BLOCKS,
        'emit': recorder.emit,
        'store': recorder.store,
        'floor_lookup': floor_of(None),
    }
    kwargs.update(overrides)
    return asyncio.run(create_document(**kwargs))


def _sheet(recorder: Recorder, **overrides) -> str:
    kwargs = {
        'request': 'REQ',
        'user': USER,
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'title': 'Daftar kontrak',
        'sheets': SHEETS,
        'emit': recorder.emit,
        'store': recorder.store,
        'floor_lookup': floor_of(None),
    }
    kwargs.update(overrides)
    return asyncio.run(create_spreadsheet(**kwargs))


def test_a_document_is_stored_for_the_asking_user_and_attached_to_the_answer():
    recorder = Recorder()
    text = _run(recorder, classification='terbatas', metadata={'owner': 'Sekretariat', 'number': ''})
    (request, rendered, user, chat_id, message_id) = recorder.stored[0]
    assert (request, user, chat_id, message_id) == ('REQ', USER, 'chat-1', 'msg-1')
    assert rendered.classification == 'TERBATAS'
    assert recorder.events == [
        {
            'type': 'files',
            'data': {
                'files': [
                    {
                        'type': 'file',
                        'id': 'file-1',
                        'name': 'JUP-NOT-Rapat-mingguan-operasi-Draf.docx',
                        'size': len(rendered.data),
                        'content_type': office_tool.DOCX_MIME,
                    }
                ]
            },
        }
    ]
    assert text.startswith('Created "JUP-NOT-Rapat-mingguan-operasi-Draf.docx" from the Notulen rapat template')
    assert 'Status: Draf. Classification: TERBATAS.' in text
    assert 'only the user who asked can download it' in text
    assert 'number, version, date, distribution' in text
    assert Document(io.BytesIO(rendered.data)).core_properties.author == 'Staf JUP'


def test_the_stored_file_records_what_it_is():
    recorder = Recorder()
    _run(recorder)
    rendered = recorder.stored[0][1]
    assert office_tool.file_metadata(rendered, 'chat-1', 'msg-1') == {
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'jup_office': {'kind': 'docx', 'template': 'NOT', 'classification': 'INTERNAL', 'status': 'Draf'},
    }


def test_a_spec_the_renderer_refuses_is_reported_and_nothing_is_stored():
    recorder = Recorder()
    text = _run(recorder, classification='PUBLIK')
    assert text.startswith(NOT_CREATED) and 'PUBLIK' in text
    assert recorder.stored == [] and recorder.events == []


def test_a_storage_failure_is_not_reported_as_a_document():
    recorder = Recorder(fail=True)
    assert _run(recorder) == STORE_FAILED
    assert recorder.events == []


def test_outside_a_chat_or_without_a_user_nothing_is_rendered():
    recorder = Recorder()
    assert _run(recorder, chat_id=None) == NO_CHAT
    assert _run(recorder, message_id='') == NO_CHAT
    assert _run(recorder, user=None) == NO_USER
    assert _run(recorder, user={'name': 'no id'}) == NO_USER
    assert recorder.stored == []


def test_a_floor_from_the_caller_raises_the_classification():
    recorder = Recorder()
    _run(recorder, classification='INTERNAL', floor='RAHASIA')
    assert recorder.stored[0][1].classification == 'RAHASIA'


def test_every_refusal_is_unmistakable():
    for message in (NO_CHAT, NO_USER, STORE_FAILED, FLOOR_FAILED):
        assert message.startswith(NOT_CREATED)


def test_a_chat_that_read_restricted_material_raises_the_documents_classification():
    recorder = Recorder()
    lookup = floor_of('TERBATAS')
    text = _run(recorder, classification='INTERNAL', floor_lookup=lookup)
    assert lookup.asked == ['chat-1']
    assert recorder.stored[0][1].classification == 'TERBATAS'
    assert 'Classification: TERBATAS' in text


def test_the_higher_of_the_callers_floor_and_the_chats_wins():
    recorder = Recorder()
    _run(recorder, floor='RAHASIA', floor_lookup=floor_of('TERBATAS'))
    assert recorder.stored[0][1].classification == 'RAHASIA'


def test_no_document_when_the_chats_floor_cannot_be_read():
    recorder = Recorder()
    assert _run(recorder, floor_lookup=broken_floor) == FLOOR_FAILED
    assert recorder.stored == [] and recorder.events == []


# -- create_spreadsheet ------------------------------------------------------------------------


def test_a_spreadsheet_is_stored_for_the_asking_user_and_attached():
    recorder = Recorder()
    text = _sheet(recorder, classification='terbatas')
    (request, rendered, user, chat_id, message_id) = recorder.stored[0]
    assert (request, user, chat_id, message_id) == ('REQ', USER, 'chat-1', 'msg-1')
    assert rendered.classification == 'TERBATAS'
    item = recorder.events[0]['data']['files'][0]
    assert item['content_type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert item['name'] == 'JUP-XLS-Daftar-kontrak-Draf.xlsx'
    assert load_workbook(io.BytesIO(rendered.data))['Kontrak']['B2'].value == 5
    assert text.startswith('Created "JUP-XLS-Daftar-kontrak-Draf.xlsx" with 1 sheet(s) and 1 row(s).')
    assert 'Status: Draf. Classification: TERBATAS.' in text


def test_the_stored_spreadsheet_records_what_it_is():
    recorder = Recorder()
    _sheet(recorder)
    rendered = recorder.stored[0][1]
    assert office_tool.file_metadata(rendered, 'chat-1', 'msg-1') == {
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'jup_office': {'kind': 'xlsx', 'classification': 'INTERNAL', 'status': 'Draf'},
    }


def test_a_chat_that_read_restricted_material_raises_the_spreadsheets_classification():
    recorder = Recorder()
    text = _sheet(recorder, classification='INTERNAL', floor_lookup=floor_of('TERBATAS'))
    assert recorder.stored[0][1].classification == 'TERBATAS'
    assert 'Classification: TERBATAS' in text


def test_spreadsheet_refusals_are_unmistakable_and_store_nothing():
    recorder = Recorder()
    assert _sheet(recorder, chat_id=None).startswith(SHEET_NOT_CREATED)
    assert _sheet(recorder, user=None).startswith(SHEET_NOT_CREATED)
    assert _sheet(recorder, floor_lookup=broken_floor).startswith(SHEET_NOT_CREATED)
    refused = _sheet(recorder, sheets=[{'name': 'X', 'columns': [{'header': 'N', 'type': 'idr'}], 'rows': [['Rp 5']]}])
    assert refused.startswith(SHEET_NOT_CREATED) and 'must be a number' in refused
    assert recorder.stored == [] and recorder.events == []
    assert _sheet(Recorder(fail=True)).startswith(SHEET_NOT_CREATED)


# -- create_presentation -----------------------------------------------------------------------


def _deck(recorder: Recorder, **overrides) -> str:
    kwargs = {
        'request': 'REQ',
        'user': USER,
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'title': 'Kinerja triwulan III',
        'slides': SLIDES,
        'emit': recorder.emit,
        'store': recorder.store,
        'floor_lookup': floor_of(None),
    }
    kwargs.update(overrides)
    return asyncio.run(create_presentation(**kwargs))


def test_a_presentation_is_stored_for_the_asking_user_and_attached():
    recorder = Recorder()
    text = _deck(recorder, cover={'kind': 'Rapat Direksi'}, classification='terbatas')
    (request, rendered, user, chat_id, message_id) = recorder.stored[0]
    assert (request, user, chat_id, message_id) == ('REQ', USER, 'chat-1', 'msg-1')
    assert rendered.classification == 'TERBATAS'
    item = recorder.events[0]['data']['files'][0]
    assert item['content_type'] == 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    assert item['name'] == 'JUP-PPT-Kinerja-triwulan-III-Draf.pptx'
    prs = Presentation(io.BytesIO(rendered.data))
    assert len(prs.slides) == 2 and prs.core_properties.author == 'Staf JUP'
    assert text.startswith('Created "JUP-PPT-Kinerja-triwulan-III-Draf.pptx" with 2 slides, including the cover')
    assert 'Status: Draf. Classification: TERBATAS.' in text


def test_the_stored_presentation_records_what_it_is():
    recorder = Recorder()
    _deck(recorder)
    rendered = recorder.stored[0][1]
    assert office_tool.file_metadata(rendered, 'chat-1', 'msg-1') == {
        'chat_id': 'chat-1',
        'message_id': 'msg-1',
        'jup_office': {'kind': 'pptx', 'classification': 'INTERNAL', 'status': 'Draf'},
    }


def test_a_chat_that_read_restricted_material_raises_the_decks_classification():
    recorder = Recorder()
    text = _deck(recorder, classification='INTERNAL', floor_lookup=floor_of('TERBATAS'))
    assert recorder.stored[0][1].classification == 'TERBATAS'
    assert 'Classification: TERBATAS' in text


def test_what_is_left_for_the_user_is_named():
    recorder = Recorder()
    text = _deck(recorder, slides=[{'layout': 'project', 'title': 'PLTS Atap', 'text': 'Tujuan'}])
    assert 'Left for the user to fill: slide 2: the project photo.' in text


def test_presentation_refusals_are_unmistakable_and_store_nothing():
    recorder = Recorder()
    assert _deck(recorder, chat_id=None).startswith(DECK_NOT_CREATED)
    assert _deck(recorder, user=None).startswith(DECK_NOT_CREATED)
    assert _deck(recorder, floor_lookup=broken_floor).startswith(DECK_NOT_CREATED)
    refused = _deck(recorder, slides=[{'layout': 'gallery', 'title': 'x'}])
    assert refused.startswith(DECK_NOT_CREATED) and 'use one of: agenda' in refused
    assert recorder.stored == [] and recorder.events == []
    assert _deck(Recorder(fail=True)).startswith(DECK_NOT_CREATED)


def test_presentation_layouts_lists_or_describes():
    assert '- kpi: Empat KPI berdampingan' in presentation_layouts()
    assert presentation_layouts('kpi').startswith('kpi: Empat KPI berdampingan')


def test_document_templates_lists_or_outlines():
    assert '- MEM: Memo usulan keputusan' in document_templates()
    assert document_templates('mem').startswith('# MEM: Memo usulan keputusan')
    assert document_templates('SRT').startswith('Unknown document template')


def test_the_stored_wrapper_holds_no_logic_of_its_own():
    """The DB copy is exec'd untested, so it may only delegate. A rule added there escapes CI."""
    source = TOOL_SOURCE.read_text()
    assert 'risk_class: write_own' in source.split('"""')[1]
    tree = ast.parse(source)
    tools = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == 'Tools')
    methods = [n for n in tools.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]
    assert [m.name for m in methods] == [
        'document_templates',
        'create_document',
        'create_spreadsheet',
        'presentation_layouts',
        'create_presentation',
    ]
    targets = (
        '_document_templates(',
        '_create_document(',
        '_create_spreadsheet(',
        '_presentation_layouts(',
        '_create_presentation(',
    )
    for method, target in zip(methods, targets):
        body = [n for n in method.body if not (isinstance(n, ast.Expr) and isinstance(n.value, ast.Constant))]
        assert len(body) == 1 and isinstance(body[0], ast.Return), method.name
        assert target in ast.unparse(body[0])


@pytest.mark.parametrize('target', ['_create_document', '_create_spreadsheet', '_create_presentation'])
def test_the_wrapper_passes_the_callers_identity_and_chat_through(target):
    call = ast.unparse(
        next(
            n
            for n in ast.walk(ast.parse(TOOL_SOURCE.read_text()))
            if isinstance(n, ast.Call) and getattr(n.func, 'id', '') == target
        )
    )
    for wire in (
        'request=__request__',
        'user=__user__',
        'chat_id=__chat_id__',
        'message_id=__message_id__',
        'emit=__event_emitter__',
    ):
        assert wire in call
    # The wrapper must not choose a floor: that comes from the chat (`jup.office_floor`).
    assert 'floor' not in call


@pytest.mark.parametrize('module', [office_tool, office_floor])
def test_the_module_loads_without_the_app(module):
    """Vendor imports live inside the I/O functions, so CI's `--only-group dev` can import the module."""
    tree = ast.parse(Path(module.__file__).read_text(encoding='utf-8'))
    top_level = {
        (node.module if isinstance(node, ast.ImportFrom) else alias.name).split('.')[0]
        for node in tree.body
        if isinstance(node, (ast.Import, ast.ImportFrom))
        for alias in node.names
    }
    assert not top_level & {'open_webui', 'fastapi', 'starlette', 'sqlalchemy'}
