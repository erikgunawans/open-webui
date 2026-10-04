"""The spreadsheet renderer (O2 1/3): `jup.office_sheet` and the chat floor in `jup.office_floor`."""

import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from jup import office_floor
from jup.office_doc import OfficeDocError
from jup.office_sheet import MAX_ROWS, MAX_SHEETS, XLSX_MIME, filename_for, render_sheet

NOW = dt.datetime(2026, 10, 3, 21, 30, tzinfo=dt.timezone(dt.timedelta(hours=7)))
CONTRACTS = {
    'name': 'Kontrak',
    'columns': [
        {'header': 'Mitra', 'type': 'text'},
        {'header': 'Nilai', 'type': 'idr'},
        {'header': 'Porsi', 'type': 'percent'},
        {'header': 'Berakhir', 'type': 'date'},
        {'header': 'Volume', 'type': 'number'},
    ],
    'rows': [
        ['PT Alpha', 1250000000, 12.5, '2027-03-31', 10.5],
        ['PT Beta', '750000000', None, '', 4],
    ],
}


def _book(rendered):
    return load_workbook(io.BytesIO(rendered.data))


def _render(sheets=None, **kwargs):
    return render_sheet('Daftar kontrak aktif', sheets or [CONTRACTS], now=NOW, **kwargs)


def test_numbers_stay_numbers_and_are_formatted():
    rendered = _render()
    ws = _book(rendered)['Kontrak']
    assert [c.value for c in ws[1]] == ['Mitra', 'Nilai', 'Porsi', 'Berakhir', 'Volume']
    assert ws['B2'].value == 1250000000 and ws['B2'].data_type == 'n'
    assert ws['B2'].number_format == '"Rp" #,##0'
    assert ws['B3'].value == 750000000
    assert ws['C2'].value == pytest.approx(0.125) and ws['C2'].number_format == '0.0%'
    assert ws['C3'].value is None
    assert ws['D2'].value == dt.datetime(2027, 3, 31) and ws['D2'].number_format == 'dd/mm/yyyy'
    assert ws['D3'].value is None
    assert ws['E2'].number_format == '#,##0.00'


def test_header_is_frozen_filtered_and_branded():
    ws = _book(_render())['Kontrak']
    assert ws.freeze_panes == 'A2'
    assert ws.auto_filter.ref == 'A1:E3'
    assert ws['A1'].font.bold and ws['A1'].font.name == 'Space Grotesk'
    assert ws['A1'].fill.fgColor.rgb.endswith('006C34')


def test_the_info_sheet_comes_first_with_status_and_classification():
    rendered = _render(classification='terbatas', author='Staf JUP')
    book = _book(rendered)
    assert book.sheetnames == ['Info', 'Kontrak']
    info = {row[0].value: row[1].value for row in book['Info'].iter_rows()}
    assert info['Status'] == 'Draf'
    assert info['Klasifikasi'] == 'TERBATAS'
    assert info['Tanggal dibuat'] == '03/10/2026 21:30 WIB'
    assert 'Staf JUP' in info['Dibuat oleh']
    assert book['Kontrak'].oddFooter.left.text == 'TERBATAS · Draf'
    assert book.properties.keywords == 'JUP-XLS'
    assert (rendered.classification, rendered.status, rendered.content_type) == ('TERBATAS', 'Draf', XLSX_MIME)
    assert rendered.filename == 'JUP-XLS-Daftar-kontrak-aktif-Draf.xlsx'


@pytest.mark.parametrize('payload', ['=SUM(A1:A9)', '+1-2', '-cmd', '@SUM(1)', '=HYPERLINK("http://x")'])
def test_nothing_the_model_writes_becomes_a_formula(payload):
    sheet = {'name': 'Data', 'columns': [{'header': 'Catatan', 'type': 'text'}], 'rows': [[payload]]}
    ws = _book(_render([sheet]))['Data']
    assert ws['A2'].value == payload
    assert ws['A2'].data_type == 's'


def test_a_header_that_looks_like_a_formula_stays_text():
    sheet = {'name': 'Data', 'columns': [{'header': '=1+1', 'type': 'text'}], 'rows': []}
    assert _book(_render([sheet]))['Data']['A1'].data_type == 's'


def test_totals_are_computed_values_not_formulas():
    ws = _book(_render([{**CONTRACTS, 'totals': True}]))['Kontrak']
    assert ws['A4'].value == 'Total'
    assert ws['B4'].value == 2000000000 and ws['B4'].data_type == 'n'
    assert ws['E4'].value == pytest.approx(14.5)
    assert ws['C4'].value is None  # percentages are not summed
    assert ws.auto_filter.ref == 'A1:E3'  # the total row stays out of the filter


@pytest.mark.parametrize('value', ['Rp 1.250.000', '1.250.000', '1,5', True, 'banyak'])
def test_amounts_that_would_need_guessing_are_refused(value):
    sheet = {'name': 'Data', 'columns': [{'header': 'Nilai', 'type': 'idr'}], 'rows': [[value]]}
    with pytest.raises(OfficeDocError, match='must be a number'):
        _render([sheet])


@pytest.mark.parametrize('value', ['03/10/2026', '2026-13-01', '3 Oktober 2026'])
def test_dates_must_be_iso(value):
    sheet = {'name': 'Data', 'columns': [{'header': 'Tanggal', 'type': 'date'}], 'rows': [[value]]}
    with pytest.raises(OfficeDocError, match='YYYY-MM-DD'):
        _render([sheet])


def test_short_rows_are_padded_and_long_rows_refused():
    ok = {'name': 'Data', 'columns': [{'header': 'A'}, {'header': 'B'}], 'rows': [['x']]}
    assert _book(_render([ok]))['Data']['B2'].value is None
    long = {'name': 'Data', 'columns': [{'header': 'A'}], 'rows': [['x', 'y']]}
    with pytest.raises(OfficeDocError, match='has 2 cells but the sheet has 1 columns'):
        _render([long])


@pytest.mark.parametrize(
    'sheets, message',
    [
        ([], 'at least one sheet'),
        ([{'name': 'Info', 'columns': [{'header': 'A'}]}], 'reserved'),
        ([{'name': 'A/B', 'columns': [{'header': 'A'}]}], 'may not contain'),
        ([{'name': 'X', 'columns': [{'header': 'A'}]}, {'name': 'x', 'columns': [{'header': 'A'}]}], 'unique'),
        ([{'name': 'X', 'columns': [{'header': 'A', 'type': 'formula'}]}], 'use one of'),
        ([{'name': 'X', 'columns': []}], 'needs columns'),
        ([{'name': 'X', 'columns': [{'header': 'A'}], 'chart': True}], 'unknown keys'),
        ([{'name': f'S{i}', 'columns': [{'header': 'A'}]} for i in range(MAX_SHEETS + 1)], 'the limit is 10'),
        ([{'name': 'X', 'columns': [{'header': 'A'}], 'rows': [[1]] * (MAX_ROWS + 1)}], 'Split it'),
    ],
)
def test_specs_outside_the_rules_are_refused_with_a_reason(sheets, message):
    with pytest.raises(OfficeDocError, match=message):
        render_sheet('Uji', sheets, now=NOW)


def test_publik_is_refused_and_the_floor_only_raises():
    with pytest.raises(OfficeDocError, match='PUBLIK'):
        _render(classification='PUBLIK')
    assert _render(classification='INTERNAL', floor='TERBATAS').classification == 'TERBATAS'
    assert _render(classification='RAHASIA', floor='TERBATAS').classification == 'RAHASIA'


def test_control_characters_are_dropped_not_fatal():
    sheet = {'name': 'Data', 'columns': [{'header': 'A'}], 'rows': [['baris\x01satu']]}
    assert _book(_render([sheet]))['Data']['A2'].value == 'barissatu'


def test_filename_is_safe():
    assert filename_for('RKA 2027 / unit: Gas & SPBG') == 'JUP-XLS-RKA-2027-unit-Gas-SPBG-Draf.xlsx'
    assert filename_for('!!!') == 'JUP-XLS-lembar-kerja-Draf.xlsx'


# -- the chat floor ----------------------------------------------------------------------------


def test_the_floor_is_read_from_the_chat_meta():
    assert office_floor.floor_from_meta(None) is None
    assert office_floor.floor_from_meta({}) is None
    assert office_floor.floor_from_meta({office_floor.KEY: {'level': 'PUBLIK'}}) is None
    assert office_floor.floor_from_meta({office_floor.KEY: {'level': 'TERBATAS'}}) == 'TERBATAS'


def test_the_floor_only_rises():
    assert office_floor.raised({}, 'TERBATAS', 'contracts', at=1) == {
        'level': 'TERBATAS',
        'source': 'contracts',
        'at': 1,
    }
    meta = {office_floor.KEY: {'level': 'TERBATAS'}}
    assert office_floor.raised(meta, 'TERBATAS', 'contracts', at=2) is None
    assert office_floor.raised(meta, 'INTERNAL', 'contracts', at=2) is None
    assert office_floor.raised(meta, 'RAHASIA', 'contracts', at=2)['level'] == 'RAHASIA'
    with pytest.raises(ValueError):
        office_floor.raised({}, 'PUBLIK', 'contracts', at=1)


def test_higher():
    assert office_floor.higher(None, None) is None
    assert office_floor.higher('INTERNAL', None) == 'INTERNAL'
    assert office_floor.higher(None, 'TERBATAS') == 'TERBATAS'
    assert office_floor.higher('RAHASIA', 'TERBATAS') == 'RAHASIA'


# -- the per-turn mark (E7.4, #88) -------------------------------------------------------------


def test_marking_lists_the_message_and_still_raises_the_floor():
    record = office_floor.marked({}, 'TERBATAS', 'contracts', at=1, message_id='m1')
    assert record == {'level': 'TERBATAS', 'source': 'contracts', 'at': 1, office_floor.MESSAGES: ['m1']}


def test_a_chat_already_at_the_floor_still_lists_a_new_message_and_keeps_its_level():
    meta = {office_floor.KEY: {'level': 'RAHASIA', 'source': 'contracts', 'at': 1, office_floor.MESSAGES: ['m1']}}
    record = office_floor.marked(meta, 'TERBATAS', 'contracts', at=2, message_id='m2')
    assert record == {'level': 'RAHASIA', 'source': 'contracts', 'at': 1, office_floor.MESSAGES: ['m1', 'm2']}
    assert office_floor.marked({office_floor.KEY: record}, 'TERBATAS', 'contracts', at=3, message_id='m2') is None


def test_a_raise_keeps_the_messages_already_listed():
    meta = {office_floor.KEY: {'level': 'TERBATAS', 'source': 'contracts', 'at': 1, office_floor.MESSAGES: ['m1']}}
    record = office_floor.marked(meta, 'RAHASIA', 'contracts', at=2, message_id='m2')
    assert record['level'] == 'RAHASIA'
    assert record[office_floor.MESSAGES] == ['m1', 'm2']


def test_without_a_message_marking_is_exactly_raising():
    assert office_floor.marked({}, 'TERBATAS', 'contracts', at=1) == office_floor.raised(
        {}, 'TERBATAS', 'contracts', at=1
    )


def test_turn_restricted_reads_only_this_turn():
    meta = {office_floor.KEY: {'level': 'TERBATAS', office_floor.MESSAGES: ['m1']}}
    assert office_floor.turn_restricted(meta, 'm1') is True
    assert office_floor.turn_restricted(meta, 'm2') is False
    assert office_floor.turn_restricted({office_floor.KEY: {'level': 'TERBATAS'}}, 'm1') is False
    assert office_floor.turn_restricted({}, 'm1') is False
    assert office_floor.turn_restricted(None, 'm1') is False
    for bad in ('x', {office_floor.KEY: []}, {office_floor.KEY: {office_floor.MESSAGES: 'm1'}}):
        with pytest.raises(office_floor.MarkUnreadable):
            office_floor.turn_restricted(bad, 'm1')
