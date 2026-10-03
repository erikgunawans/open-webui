"""The slides renderer (O3 1/3): `jup.office_slides` on Template presentasi JUP v1.5-r2."""

import datetime as dt
import hashlib
import io
import json

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from jup.office_doc import OfficeDocError
from jup.office_slides import (
    MAX_SLIDES,
    PPTX_MIME,
    capacities,
    filename_for,
    layout_guide,
    lines_needed,
    manifest,
    render_slides,
    template_path,
)

NOW = dt.datetime(2026, 10, 3, 21, 30, tzinfo=dt.timezone(dt.timedelta(hours=7)))
KPIS = [{'name': f'KPI {i}', 'value': '12,5', 'unit': 'Rp miliar', 'period': 'Periode: Sep 2026'} for i in range(1, 5)]
TABLE = {
    'columns': [
        {'header': 'Uraian', 'type': 'text'},
        {'header': 'Realisasi', 'type': 'idr'},
        {'header': 'Capaian', 'type': 'percent'},
        {'header': 'Tanggal', 'type': 'date'},
    ],
    'rows': [['Pendapatan', 412000, 94.1, '2026-09-30'], ['Laba', 58000.5, None, '']],
}
CHART = {'categories': ['TW I', 'TW II', 'TW III'], 'series': [{'name': 'Realisasi', 'values': [120, 138, 154]}]}
SOURCE = 'Sumber: laporan manajemen. Periode: 2026.'

# One slide on every working layout, in manifest order.
EVERY_LAYOUT = [
    {'layout': 'agenda', 'title': 'Agenda', 'items': ['Kinerja', 'RKA 2027', 'Keputusan']},
    {'layout': 'summary', 'title': 'Ringkasan', 'points': [{'heading': 'Kondisi', 'text': 'Baik'}] * 2},
    {'layout': 'section', 'title': 'Kinerja'},
    {'layout': 'text', 'title': 'Pesan', 'message': 'Proyek gas tertunda'},
    {
        'layout': 'two_columns',
        'title': 'Dua sisi',
        'left_heading': 'Peluang',
        'left_points': ['A', 'B'],
        'right_heading': 'Tantangan',
        'right_points': 'C',
    },
    {'layout': 'kpi', 'title': 'Kinerja utama', 'kpis': KPIS, 'source': SOURCE},
    {'layout': 'bar_chart', 'title': 'Realisasi dan target', 'chart': CHART, 'source': SOURCE},
    {'layout': 'line_chart', 'title': 'Tren', 'chart': CHART, 'source': SOURCE},
    {'layout': 'financial_table', 'title': 'Keuangan', 'table': TABLE, 'source': SOURCE},
    {'layout': 'portfolio', 'title': 'Portofolio', 'solar': 'PLTS 12 MWp'},
    {'layout': 'project', 'title': 'PLTS Atap', 'text': ['Tujuan', 'Lingkup']},
    {'layout': 'schedule', 'title': 'Jadwal', 'table': TABLE},
    {
        'layout': 'process',
        'title': 'Alur',
        'steps': [{'name': f'Tahap {i}', 'input': 'Unit', 'output': 'Hasil'} for i in range(1, 5)],
    },
    {'layout': 'risks', 'title': 'Risiko', 'table': TABLE},
    {'layout': 'options', 'title': 'Opsi', 'table': TABLE},
    {
        'layout': 'decision',
        'title': 'Usulan keputusan',
        'decision': 'Menyetujui RKA 2027',
        'basis': 'Laba naik',
        'scope': 'Capex 2027',
        'follow_up': 'Keuangan',
    },
    {'layout': 'actions', 'title': 'Tindak lanjut', 'table': TABLE, 'notes': 'Tekankan tenggat.'},
    {'layout': 'closing', 'title': 'Terima kasih', 'contact': ['Sekretariat', 'sekper@jup.co.id']},
]


def _render(slides=None, **kwargs):
    kwargs.setdefault('now', NOW)
    return render_slides('Laporan kinerja triwulan III', EVERY_LAYOUT if slides is None else slides, **kwargs)


def _deck(rendered):
    return Presentation(io.BytesIO(rendered.data))


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def _all_text(prs) -> str:
    return '\n'.join(t for slide in prs.slides for t in _texts(slide))


def _with(layout: str, **fields) -> dict:
    slide = dict(next(s for s in EVERY_LAYOUT if s['layout'] == layout))
    slide.update(fields)
    return slide


# -- the template ------------------------------------------------------------------------------


def test_the_template_is_the_one_jup_issued():
    assert hashlib.sha256(template_path().read_bytes()).hexdigest() == manifest()['sha256']


def test_the_manifest_names_twenty_working_layouts_and_excludes_the_rest():
    names = [c['layout'] for c in manifest()['covers'].values()] + [s['layout'] for s in manifest()['layouts'].values()]
    assert sorted(n[:2] for n in names) == [f'{i:02d}' for i in range(1, 21)]
    assert sorted(n[:2] for n in manifest()['excluded']) == [f'{i:02d}' for i in range(21, 31)]


def test_every_manifest_placeholder_exists_on_its_layout():
    caps = capacities()
    for key, spec in manifest()['layouts'].items():
        boxes = caps[spec['layout']]
        idxs = [f['idx'] for f in spec['fields'].values()]
        idxs += [i for g in spec.get('groups', {}).values() for slot in g['slots'] for i in slot.values()]
        idxs += spec.get('fixed', []) + ([spec['source']['idx']] if spec.get('source') else [])
        for idx in idxs:
            assert idx in boxes, (key, idx)


def test_the_renderer_produces_valid_ooxml_placeholder_types():
    """The template's `subtitle` is corrected to `subTitle`, which python-pptx (and OOXML) read."""
    prs = _deck(_render())
    for slide in prs.slides:
        for shape in slide.placeholders:
            assert shape.placeholder_format.type is not None


# -- the deck ----------------------------------------------------------------------------------


def test_a_deck_has_the_cover_then_one_slide_per_spec_on_the_named_layouts():
    rendered = _render(cover={'kind': 'Rapat Direksi', 'unit': 'Divisi Keuangan', 'date': '3 Oktober 2026'})
    prs = _deck(rendered)
    layouts = manifest()['layouts']
    assert [s.slide_layout.name for s in prs.slides] == ['01 Sampul institusional'] + [
        layouts[s['layout']]['layout'] for s in EVERY_LAYOUT
    ]
    assert rendered.slides == len(EVERY_LAYOUT) + 1
    assert rendered.charts == 2 and rendered.tables == 5
    cover = _texts(prs.slides[0])
    assert 'Laporan kinerja triwulan III' in cover
    assert 'RAPAT DIREKSI' in cover
    assert 'Divisi Keuangan   3 Oktober 2026' in cover
    assert 'PT Jakarta Utilitas Propertindo' in cover


def test_the_dark_cover_is_chosen_by_style():
    prs = _deck(_render(cover={'style': 'korporat'}))
    assert prs.slides[0].slide_layout.name == '02 Sampul korporat'
    assert '03/10/2026' in _texts(prs.slides[0])  # no date given: today's, in Jakarta


def test_no_example_content_survives():
    text = _all_text(_deck(_render()))
    for leak in ('Data ilustratif', 'CONTOH DATA', '[Unit kerja]', '[Klasifikasi', 'Template presentasi', 'XX,X'):
        assert leak not in text


def test_every_shape_on_a_generated_slide_comes_from_its_layout():
    """No free-floating shapes: every text shape is a placeholder, and the rest are its table or chart."""
    for slide in _deck(_render()).slides:
        for shape in slide.shapes:
            assert shape.is_placeholder, (slide.slide_layout.name, shape.name)


def test_classification_and_draf_are_in_every_footer():
    prs = _deck(_render(classification='TERBATAS', cover={'unit': 'Divisi Keuangan'}))
    with_footer = 0
    for slide in prs.slides:
        footers = [s.text_frame.text for s in slide.placeholders if s.placeholder_format.idx == 98]
        layout_has_footer = any(s.placeholder_format.idx == 98 for s in slide.slide_layout.placeholders)
        assert bool(footers) == layout_has_footer, slide.slide_layout.name
        for footer in footers:
            assert footer == 'Divisi Keuangan / TERBATAS · Draf'
            with_footer += 1
    assert with_footer == len(prs.slides) - 2  # the section divider and the closing have no footer field


def test_the_classification_defaults_to_internal_rises_to_the_floor_and_is_never_publik():
    assert _render().classification == 'INTERNAL'
    assert _render(classification='INTERNAL', floor='TERBATAS').classification == 'TERBATAS'
    assert _render(classification='RAHASIA', floor='TERBATAS').classification == 'RAHASIA'
    with pytest.raises(OfficeDocError, match='PUBLIK'):
        _render(classification='PUBLIK')


def test_the_file_says_what_it_is():
    rendered = _render()
    assert rendered.filename == 'JUP-PPT-Laporan-kinerja-triwulan-III-Draf.pptx'
    assert rendered.content_type == PPTX_MIME and rendered.status == 'Draf'
    props = _deck(rendered).core_properties
    assert props.title == 'Laporan kinerja triwulan III'
    assert props.keywords == 'JUP-PPT' and props.subject == 'Draf · INTERNAL'
    assert filename_for('!!!') == 'JUP-PPT-presentasi-Draf.pptx'


def test_slide_numbers_are_live_fields():
    prs = _deck(_render())
    agenda = prs.slides[1]
    fields = [f.get('type') for f in agenda.shapes._spTree.iter(qn('a:fld'))]
    assert 'slidenum' in fields


def test_groups_are_numbered_and_unused_slots_removed():
    prs = _deck(_render())
    agenda = _texts(prs.slides[1])
    assert ['01', '02', '03'] == [t for t in agenda if t in ('01', '02', '03', '04')]
    assert 'Kinerja' in agenda and 'Keputusan' in agenda
    assert not any(s.placeholder_format.idx in (9, 10) for s in prs.slides[1].placeholders)


def test_section_dividers_count_themselves():
    deck = [_with('section', title='Satu'), _with('section', title='Dua')]
    prs = _deck(_render(deck))
    assert '01' in _texts(prs.slides[1]) and '02' in _texts(prs.slides[2])


def test_the_section_label_defaults_to_the_layouts_own_and_can_be_set():
    prs = _deck(_render([_with('agenda'), _with('bar_chart'), _with('agenda', section='Rapat khusus')]))
    labels = [
        next(s.text_frame.text for s in slide.placeholders if s.placeholder_format.idx == 2)
        for slide in list(prs.slides)[1:]
    ]
    assert labels == ['PEMBAHASAN', 'KINERJA', 'RAPAT KHUSUS']


def test_text_takes_the_layouts_formatting():
    """The template's sizes live on its prompt runs, not in the list style; they must be copied."""
    prs = _deck(_render())
    summary = prs.slides[2]
    title = next(s for s in summary.placeholders if s.placeholder_format.idx == 0)
    run = title.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 33 and run.font.bold and run.font.name == 'Space Grotesk'


def test_double_asterisks_make_bold_runs():
    prs = _deck(_render([_with('text', side_text='Kontrak **EPC** paling lambat Desember')]))
    side = next(s for s in prs.slides[1].placeholders if s.placeholder_format.idx == 5)
    runs = side.text_frame.paragraphs[0].runs
    assert [(r.text, bool(r.font.bold)) for r in runs] == [
        ('Kontrak ', False),
        ('EPC', True),
        (' paling lambat Desember', False),
    ]


def test_spaced_lists_put_a_blank_line_between_items():
    prs = _deck(_render([_with('two_columns')]))
    left = next(s for s in prs.slides[1].placeholders if s.placeholder_format.idx == 4)
    assert [p.text for p in left.text_frame.paragraphs] == ['A', '', 'B']


def test_speaker_notes_are_kept():
    prs = _deck(_render())
    assert prs.slides[len(EVERY_LAYOUT) - 1].notes_slide.notes_text_frame.text == 'Tekankan tenggat.'


def test_the_project_photo_is_left_for_the_user():
    rendered = _render([_with('project')])
    assert rendered.unfilled == ['slide 2: the project photo']


# -- tables ------------------------------------------------------------------------------------


def test_a_table_is_typed_formatted_and_styled_as_the_template():
    prs = _deck(_render([_with('financial_table')]))
    table = next(s for s in prs.slides[1].shapes if s.has_table).table
    assert [c.text for c in table.rows[0].cells] == ['URAIAN', 'REALISASI', 'CAPAIAN', 'TANGGAL']
    assert [c.text for c in table.rows[1].cells] == ['Pendapatan', 'Rp 412.000,0', '94,1%', '30/09/2026']
    assert [c.text for c in table.rows[2].cells] == ['Laba', 'Rp 58.000,5', '', '']
    assert table.rows[1].cells[1].text_frame.paragraphs[0].alignment == 3  # right (PP_ALIGN.RIGHT)
    assert table.rows[1].cells[0].text_frame.paragraphs[0].alignment is None
    header_run = table.rows[0].cells[0].text_frame.paragraphs[0].runs[0]
    assert str(header_run.font.color.rgb) == '004999'


def test_table_rows_keep_the_examples_height_and_are_capped_per_layout():
    """The layout draws a rule under each example row; a taller table would cut through them."""
    rows = [['x', 1, 1, '2026-01-01']] * 4
    _render([_with('financial_table', table={**TABLE, 'rows': rows})])
    with pytest.raises(OfficeDocError, match='this layout holds 3 under the header'):
        _render([_with('risks', table={**TABLE, 'rows': rows})])


def test_a_table_cell_that_wraps_is_refused():
    long = 'Koordinasi dengan dinas terkait untuk percepatan perizinan lahan'
    with pytest.raises(OfficeDocError, match='table cell'):
        _render([_with('risks', table={**TABLE, 'rows': [[long, 1, 1, '2026-01-01']]})])


@pytest.mark.parametrize(
    'bad, message',
    [
        ({'columns': [{'header': 'A'}], 'rows': [['x']]}, '2 to 6 columns'),
        ({**TABLE, 'rows': [['x', 'Rp 5', 1, '2026-01-01']]}, 'must be a number'),
        ({**TABLE, 'rows': [['x', 5, 1, '30/09/2026']]}, 'YYYY-MM-DD'),
        ({**TABLE, 'rows': []}, 'at least one row'),
        ({**TABLE, 'style': 'x'}, 'unknown keys'),
    ],
)
def test_bad_tables_are_refused(bad, message):
    with pytest.raises(OfficeDocError, match=message):
        _render([_with('financial_table', table=bad)])


# -- charts ------------------------------------------------------------------------------------


def test_charts_are_native_and_carry_the_data():
    prs = _deck(_render())
    bar = next(s for s in prs.slides[7].shapes if s.has_chart).chart
    line = next(s for s in prs.slides[8].shapes if s.has_chart).chart
    assert bar.chart_type == 51 and line.chart_type == 65  # COLUMN_CLUSTERED, LINE_MARKERS
    assert list(bar.plots[0].categories) == ['TW I', 'TW II', 'TW III']
    assert [list(s.values) for s in bar.series] == [[120.0, 138.0, 154.0]]
    # Editable in PowerPoint: the data lives in an embedded workbook.
    assert bar.part.chart_workbook.xlsx_part is not None


def test_the_value_axis_follows_the_data_not_the_example():
    """The example fixes its axis at 0-150; 154 would be cut off."""
    prs = _deck(_render())
    space = next(s for s in prs.slides[7].shapes if s.has_chart).chart._chartSpace
    assert not list(space.iter(qn('c:max'))) and not list(space.iter(qn('c:majorUnit')))
    negative = {**CHART, 'series': [{'name': 'Laba', 'values': [-5, 3, 4]}]}
    space = next(
        s for s in _deck(_render([_with('bar_chart', chart=negative)])).slides[1].shapes if s.has_chart
    ).chart._chartSpace
    assert not list(space.iter(qn('c:min')))


def test_extra_series_get_their_own_colour_and_a_legend():
    chart = {
        'categories': ['A', 'B'],
        'series': [{'name': n, 'values': [1, 2]} for n in ('Realisasi', 'Target', 'Tahun lalu')],
    }
    prs = _deck(_render([_with('line_chart', chart=chart)]))
    native = next(s for s in prs.slides[1].shapes if s.has_chart).chart
    assert native.has_legend
    colours = [str(s.format.line.color.rgb) for s in list(native.series)[1:]]
    assert len(set(colours)) == 2


def test_the_example_charts_alt_text_is_replaced():
    prs = _deck(_render())
    frame = next(s for s in prs.slides[7].shapes if s.has_chart)
    assert 'ilustratif' not in frame._element.find('.//' + qn('p:cNvPr')).get('descr')


@pytest.mark.parametrize(
    'bad, message',
    [
        ({**CHART, 'categories': ['A']}, '2 to 12 categories'),
        ({**CHART, 'series': []}, '1 to 3 series'),
        ({**CHART, 'series': [{'name': 'X', 'values': [1, 2]}]}, 'exactly 3 values'),
        ({**CHART, 'series': [{'name': 'X', 'values': [1, 'Rp 2', 3]}]}, 'must be a number'),
    ],
)
def test_bad_charts_are_refused(bad, message):
    with pytest.raises(OfficeDocError, match=message):
        _render([_with('bar_chart', chart=bad)])


# -- measuring and refusing ---------------------------------------------------------------------


def test_lines_needed_wraps_by_words():
    assert lines_needed('satu dua tiga', 8) == 2
    assert lines_needed('satu\ndua', 40) == 2
    assert lines_needed('a' * 25, 10) == 3
    assert lines_needed('**tebal** biasa', 12) == 1


def test_text_that_does_not_fit_is_refused_never_shrunk():
    box = capacities()['04 Ringkasan eksekutif'][5]
    too_long = ' '.join(['pendapatan'] * (box.chars * (box.lines + 1) // 10))
    with pytest.raises(OfficeDocError, match='does not fit.*never made smaller'):
        _render([_with('summary', points=[{'heading': 'A', 'text': too_long}] * 2)])


def test_a_title_is_one_line_when_a_subtitle_sits_under_it():
    title = 'Realisasi pendapatan dan beban operasi triwulan III'
    _render([_with('summary', title=title)])
    with pytest.raises(OfficeDocError, match='one line'):
        _render([_with('summary', title=title, subtitle='Pendapatan tumbuh.')])


@pytest.mark.parametrize(
    'slide, message',
    [
        ({'layout': 'gallery', 'title': 'x'}, 'use one of: agenda'),
        ({'layout': 'agenda', 'items': ['a', 'b']}, 'needs "title"'),
        (_with('agenda', items=['a']), '2 to 4 items'),
        (_with('kpi', kpis=KPIS[:3]), 'exactly 4 items'),
        (_with('kpi', source=None), 'needs "source"'),
        (_with('bar_chart', source=''), 'needs "source"'),
        (_with('agenda', colour='red'), 'unknown keys: colour'),
        (_with('section', section='X'), 'unknown keys: section'),
        (_with('portfolio', solar=None), 'at least one of'),
        (_with('summary', points=[{'heading': 'A'}] * 2), 'needs "text"'),
        (_with('summary', points=[{'heading': 'A', 'text': 'b', 'icon': 'x'}] * 2), 'unknown keys: icon'),
    ],
)
def test_bad_slides_are_refused_with_a_message_the_model_can_act_on(slide, message):
    with pytest.raises(OfficeDocError, match=message):
        _render([slide])


def test_limits_are_refused_not_truncated():
    with pytest.raises(OfficeDocError, match=f'the limit is {MAX_SLIDES}'):
        _render([_with('agenda')] * MAX_SLIDES)
    with pytest.raises(OfficeDocError, match='at least one slide'):
        _render([])
    with pytest.raises(OfficeDocError, match='cover style'):
        _render(cover={'style': 'neon'})
    with pytest.raises(OfficeDocError, match='unknown keys'):
        _render(cover={'logo': 'x'})


# -- the model's guide --------------------------------------------------------------------------


def test_the_guide_lists_layouts_and_describes_one():
    listing = layout_guide()
    assert all(f'- {key}:' in listing for key in manifest()['layouts'])
    kpi = layout_guide('KPI')
    assert 'kpis: exactly 4 objects' in kpi and 'source (required' in kpi
    assert 'characters' in kpi
    assert '1 to 3 rows' in layout_guide('risks')
    assert layout_guide('gallery').startswith('Unknown layout')


def test_the_manifest_is_valid_json_with_tabs():
    raw = (template_path().parent / 'manifest.json').read_text(encoding='utf-8')
    assert json.loads(raw)['file'] == template_path().name


def test_shape_ids_are_unique_on_every_slide():
    """PowerPoint offers to "repair" a file whose slide has two shapes with one id; LibreOffice does not care."""
    for slide in _deck(_render()).slides:
        ids = [e.get('id') for e in slide.shapes._spTree.iter(qn('p:cNvPr'))]
        assert len(ids) == len(set(ids)), slide.slide_layout.name
